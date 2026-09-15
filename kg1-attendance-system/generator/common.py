# -*- coding: utf-8 -*-
"""Shared constants, styles and helpers for the KG1 attendance workbook generator."""
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName

FONT = "Arial"

# ---------- palette ----------
NAVY = "1F3A5F"; TEAL = "2A7F8E"; TEAL_L = "E3F0F2"; LIGHT = "F5F7FA"; BORDER = "D9DEE5"
TXT = "1F2933"; MUTED = "6B7280"; WHITE = "FFFFFF"; INPUT_FILL = "FFF8E1"; INPUT_BORDER = "E6C86E"
GREEN_F, GREEN_T = "E8F5E9", "2E7D32"
YEL_F, YEL_T = "FFF4CC", "8A6D00"
ORG_F, ORG_T = "FFE0B2", "B45309"
DORG_F, DORG_T = "FDD0C0", "C2410C"
RED_F, RED_T = "F8B4B4", "B91C1C"
ABS_F, ABS_T = "FADBD8", "B03A2E"
EXC_F, EXC_T = "FFF3CD", "856404"
LATE_F, LATE_T = "DCE8F5", "1F4E79"
GRAY_F = "EDEDED"; GRAY_F2 = "F3F4F6"; BLUE_F = "DBEAFE"

# ---------- academic calendar (columns are generated once; dates are static) ----------
YEAR_START = dt.date(2026, 8, 31)   # Monday
YEAR_END = dt.date(2027, 6, 30)
DATES = []
_d = YEAR_START
while _d <= YEAR_END:
    if _d.weekday() < 5:          # Monday..Friday (UAE school week)
        DATES.append(_d)
    _d += dt.timedelta(days=1)
NDAYS = len(DATES)                 # 218
SLOTS = 30                         # student capacity per class sheet
CLASSES = ["KG1-1", "KG1-2", "KG1-3", "KG1-4", "KG1-5", "KG1-6", "KG1-7"]
ACTIVE_CLASSES = 7
NCLS = len(CLASSES)

# register layout
REG_FIRST = 9                      # first student row
REG_LAST = REG_FIRST + SLOTS - 1   # 38
DATE_COL0 = 30                     # column AD (after the hidden helper block)
DATE_COL1 = DATE_COL0 + NDAYS - 1  # 247 -> IM
DC0 = L(DATE_COL0); DC1 = L(DATE_COL1)
SUM_ROW = {"abs": 40, "exc": 41, "late": 42, "pres": 43, "reg": 44, "rate": 45}

# master layout
STU_FIRST = 5
STU_LAST = STU_FIRST + NCLS * SLOTS - 1   # 214
# database layout
DB_FIRST = 2
DB_BLOCK = SLOTS * NDAYS                  # 6540 rows per class
DB_LAST = DB_FIRST + NCLS * DB_BLOCK - 1  # 45781
# calendar layout
CAL_FIRST = 5
CAL_LAST = CAL_FIRST + NDAYS - 1          # 222
# contact log
LOG_FIRST = 6; LOG_LAST = 305
# absence log
ALOG_FIRST = 6; ALOG_LAST = 1505

AR_DAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]
AR_MONTHS = ["يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو", "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر"]

SHEETS = {
    "dash": "لوحة المتابعة", "daily": "الحضور اليومي", "fu": "قائمة المتابعة", "alog": "سجل الغياب",
    "clog": "سجل التواصل", "prof": "ملف الطفل", "trend": "الاتجاهات", "rpt": "التقرير",
    "stu": "الطلاب", "cal": "التقويم", "set": "الإعدادات", "guide": "الدليل", "db": "قاعدة البيانات", "calc": "حسابات",
}

def q(sheet):
    return "'" + sheet.replace("'", "''") + "'"

def cq(sheet):  # sheet ref prefix for formulas
    return q(sheet) + "!"

# ---------- styles ----------
def font(size=10, bold=False, color=TXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)

def side(color=BORDER, style="thin"):
    return Side(style=style, color=color)

def box(color=BORDER):
    s = side(color)
    return Border(left=s, right=s, top=s, bottom=s)

def bottom(color=BORDER, style="thin"):
    return Border(bottom=side(color, style))

ALIGN_R = Alignment(horizontal="right", vertical="center", wrap_text=False)
ALIGN_C = Alignment(horizontal="center", vertical="center")
ALIGN_CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_RW = Alignment(horizontal="right", vertical="center", wrap_text=True)
UNLOCKED = Protection(locked=False)

def style(c, *, f=None, bg=None, al=None, b=None, nf=None, unlock=False):
    if f is not None: c.font = f
    if bg is not None: c.fill = fill(bg)
    if al is not None: c.alignment = al
    if b is not None: c.border = b
    if nf is not None: c.number_format = nf
    if unlock: c.protection = UNLOCKED
    return c

def put(ws, ref, value=None, **kw):
    c = ws[ref] if isinstance(ref, str) else ws.cell(*ref)
    if value is not None:
        c.value = value
    return style(c, **kw)

def header_row(ws, row, col0, labels, *, bg=NAVY, fg=WHITE, size=10, height=30, widths=None, wrap=True):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col0 + i, lab)
        style(c, f=font(size, True, fg), bg=bg, al=ALIGN_CW if wrap else ALIGN_C, b=box(bg))
        if widths:
            ws.column_dimensions[L(col0 + i)].width = widths[i]
    ws.row_dimensions[row].height = height

def title_block(ws, title, subtitle=None, en=None, row=1):
    put(ws, (row, 1), title, f=font(18, True, NAVY), al=ALIGN_R)
    ws.row_dimensions[row].height = 30
    if subtitle:
        put(ws, (row + 1, 1), subtitle, f=font(11, False, TEAL), al=ALIGN_R)
    if en:
        put(ws, (row + 2, 1), en, f=font(9, False, MUTED, italic=True), al=ALIGN_R)

def section(ws, ref, ar, en=None, width_cols=1):
    c = put(ws, ref, ar + (("   |   " + en) if en else ""), f=font(12, True, NAVY), al=ALIGN_R)
    c.border = Border(bottom=side(TEAL, "medium"))
    # extend the underline across the section width
    r, col = c.row, c.column
    for i in range(1, width_cols):
        ws.cell(r, col + i).border = Border(bottom=side(TEAL, "medium"))

def input_cell(ws, ref, value=None, nf=None, al=ALIGN_C):
    return put(ws, ref, value, f=font(10, True, NAVY), bg=INPUT_FILL, al=al, b=box(INPUT_BORDER), nf=nf, unlock=True)

def setup_sheet(ws, rtl=True, grid=False, zoom=100, tab=None):
    ws.sheet_view.rightToLeft = rtl
    ws.sheet_view.showGridLines = grid
    ws.sheet_view.zoomScale = zoom
    if tab:
        ws.sheet_properties.tabColor = tab

def define(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

def rng(sheet, c0, r0, c1, r1, abs_=True):
    d = "$" if abs_ else ""
    return f"{cq(sheet)}{d}{L(c0)}{d}{r0}:{d}{L(c1)}{d}{r1}"

def protect(ws, password=None):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.sort = False
    p.autoFilter = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    if password:
        p.password = password


APPROVER = "نعيمة الكتبي"
PRINCIPAL = "منى الغيثي"
import os
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")

def add_logo(ws, name, c0, r0, c1, r1):
    """Place a logo PNG inside a cell rectangle (1-based, inclusive) - works in RTL sheets."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
    path = os.path.join(ASSETS, name)
    if not os.path.exists(path):
        return
    img = XLImage(path)
    a = TwoCellAnchor(editAs="oneCell")
    a._from = AnchorMarker(col=c0 - 1, colOff=40000, row=r0 - 1, rowOff=20000)
    a.to = AnchorMarker(col=c1, colOff=-40000, row=r1, rowOff=-20000)
    img.anchor = a
    ws.add_image(img)

def approval_block(ws, label_ref, name_ref, date_ref, size=10):
    put(ws, label_ref, "يعتمد،", f=font(size, True, NAVY), al=ALIGN_R)
    put(ws, name_ref, APPROVER, f=font(size + 1, True, NAVY), al=ALIGN_R)
    put(ws, date_ref, '="التاريخ: "&TEXT(TODAY(),"dd/mm/yyyy")', f=font(size - 1, False, MUTED), al=ALIGN_R)

def principal_block(ws, label_ref, name_ref, sig_ref, size=10):
    AL = Alignment(horizontal="left", vertical="center")
    put(ws, label_ref, "مديرة المدرسة:", f=font(size, True, NAVY), al=AL)
    put(ws, name_ref, PRINCIPAL, f=font(size + 1, True, NAVY), al=AL)
    put(ws, sig_ref, "التوقيع: ______________________", f=font(size - 1, False, MUTED), al=AL)

def letterhead(ws, logo_cols, rows=(1, 3)):
    """Ministry logo centered, school name at the far right (column A) with the zone beneath it."""
    add_logo(ws, "moe_logo.png", logo_cols[0], rows[0], logo_cols[1], rows[1])
    put(ws, "A1", "=SchoolName", f=font(16, True, NAVY), al=Alignment(horizontal="right", vertical="center"))
    put(ws, "A2", '="نطاق "&SchoolZone', f=font(11, True, TEAL), al=Alignment(horizontal="right", vertical="center"))

# register columns (visible block A..L, hidden helper block M..AC)
RC = {"num": "A", "id": "B", "name": "C", "abs": "D", "exc": "E", "late": "F", "rate": "G", "status": "H", "contact": "I",
      "phone": "J", "guardian": "K", "notes": "L", "days": "M", "pres": "N", "absrate": "O", "lastabs": "P", "contacted": "Q",
      "lastcontact": "R", "method": "S", "level": "T", "start": "U", "d3": "V", "d5": "W", "d10": "X", "d15": "Y",
      "ph1": "Z", "ph2": "AA", "ph3": "AB", "ph4": "AC"}
CUTOFF_CELL = "$L$2"      # effective "counted up to" date inside each register sheet

# ---- cross-sheet helpers (no background database): pick a class grid / counted-row by class index
def GRID(c_expr):
    return "CHOOSE(" + c_expr + "," + ",".join(f"{cq(cn)}${DC0}${REG_FIRST}:${DC1}${REG_LAST}" for cn in CLASSES) + ")"

def CONFROW(c_expr):
    return "CHOOSE(" + c_expr + "," + ",".join(f"{cq(cn)}${DC0}$7:${DC1}$7" for cn in CLASSES) + ")"

DATES_ROW = f"{cq(CLASSES[0])}${DC0}$5:${DC1}$5"

def CODE_OF(val_expr):
    """Arabic status cell -> P/A/E/L (blank cell on a counted day = present)."""
    return f'IF({val_expr}="","P",IFERROR(INDEX(StatusCodes,MATCH({val_expr},StatusLabels,0)),""))'

SHEETS["cd"] = "ملخص_الصفوف"
SHEETS["ev"] = "أحداث_الغياب"
EV_PER_STUDENT = 20
