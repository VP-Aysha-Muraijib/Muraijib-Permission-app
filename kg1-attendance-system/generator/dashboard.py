# -*- coding: utf-8 -*-
from common import *
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import DoughnutChart, BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.label import DataLabelList
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker

def place(ws, chart, c0, r0, c1, r1):
    """Anchor a chart to a cell rectangle (1-based, inclusive) so it renders identically in RTL sheets."""
    a = TwoCellAnchor(editAs="oneCell")
    a._from = AnchorMarker(col=c0 - 1, colOff=0, row=r0 - 1, rowOff=0)
    a.to = AnchorMarker(col=c1, colOff=0, row=r1, rowOff=0)
    chart.anchor = a
    ws.add_chart(chart)

def labels(chart, fmt="0%"):
    d = DataLabelList(); d.showVal = True; d.showSerName = False; d.showCatName = False; d.showLegendKey = False; d.showPercent = False; d.numFmt = fmt
    chart.dataLabels = d

CALC = SHEETS["calc"]; DASH = SHEETS["dash"]

def class_counts(cls_expr, code, date_expr=None, start=None, end=None):
    """COUNTIFS over the database, with optional class filter ('الكل' = all)."""
    crit = f'DB_Code,"{code}"'
    if date_expr: crit += f',DB_Date,{date_expr}'
    if start: crit += f',DB_Date,">="&{start},DB_Date,"<="&{end}'
    return f'IF({cls_expr}="الكل",COUNTIFS({crit}),COUNTIFS({crit},DB_Class,{cls_expr}))'

def build_calc(wb):
    ws = wb.create_sheet(CALC)
    setup_sheet(ws, grid=True)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 16
    put(ws, "A1", "حسابات مساعدة للوحة المتابعة (تلقائي – لا تعدّل)", f=font(11, True, NAVY))
    cells = {
        "LastRecorded": ('آخر يوم مسجَّل', f'=IF(COUNTIF(Cal_Reg,">0")=0,YearStart,_xlfn.MAXIFS(Cal_Date,Cal_Reg,">0"))', "dd/mm/yyyy"),
        "Dash_Date": ("التاريخ الفعّال", '=IF(Dash_DateInput="",LastRecorded,Dash_DateInput)', "dd/mm/yyyy"),
        "Dash_DayIdx": ("ترتيب اليوم في التقويم", '=IFERROR(MATCH(Dash_Date,Cal_Date,0),0)', None),
        "Dash_PStart": ("بداية الفترة", '=IF(Dash_Period="اليوم",Dash_Date,IF(Dash_Period="الأسبوع",Dash_Date-WEEKDAY(Dash_Date,2)+1,IF(Dash_Period="الشهر",DATE(YEAR(Dash_Date),MONTH(Dash_Date),1),IF(Dash_Period="الفصل الدراسي",IFERROR(INDEX(TermStart,MATCH(Dash_Date,TermStart,1)),YearStart),YearStart))))', "dd/mm/yyyy"),
        "Dash_PEnd": ("نهاية الفترة", '=IF(Dash_Period="اليوم",Dash_Date,IF(Dash_Period="الأسبوع",Dash_PStart+6,IF(Dash_Period="الشهر",EOMONTH(Dash_Date,0),IF(Dash_Period="الفصل الدراسي",IFERROR(INDEX(TermEnd,MATCH(Dash_Date,TermStart,1)),YearEnd),YearEnd))))', "dd/mm/yyyy"),
        "Dash_PeriodText": ("وصف الفترة", '=Dash_Period&": "&TEXT(Dash_PStart,"dd/mm/yyyy")&" – "&TEXT(Dash_PEnd,"dd/mm/yyyy")', None),
        "Dash_MStart": ("بداية الشهر", '=DATE(YEAR(Dash_Date),MONTH(Dash_Date),1)', "dd/mm/yyyy"),
        "Dash_MEnd": ("نهاية الشهر", '=EOMONTH(Dash_Date,0)', "dd/mm/yyyy"),
        "Dash_Rate": ("نسبة الحضور للفترة", '=IF(SUM(Stu_PTot)=0,"",SUM(Stu_PAtt)/SUM(Stu_PTot))', "0.0%"),
        "Today_P": ("حاضر اليوم", '=COUNTIFS(Stu_InFilter,1,Stu_Today,"P")', None),
        "Today_A": ("غائب اليوم", '=COUNTIFS(Stu_InFilter,1,Stu_Today,"A")', None),
        "Today_E": ("بعذر اليوم", '=COUNTIFS(Stu_InFilter,1,Stu_Today,"E")', None),
        "Today_L": ("متأخر اليوم", '=COUNTIFS(Stu_InFilter,1,Stu_Today,"L")', None),
        "Today_Reg": ("المسجَّلون اليوم", '=Today_P+Today_A+Today_E+Today_L', None),
        "Today_Att": ("الحاضرون فعليًا (حاضر + متأخر)", '=Today_P+Today_L', None),
        "Today_Rate": ("نسبة الحضور اليوم", '=IF(Today_Reg=0,"",(Today_P+IF(LateAsPresent="نعم",Today_L,0))/Today_Reg)', "0.0%"),
        "Total_Kids": ("إجمالي الأطفال (حسب الفلتر)", '=SUM(Stu_InFilter)', None),
        "Rec_Days": ("عدد الأيام المسجَّلة", '=COUNTIF(Cal_Reg,">0")', None),
        "Cls_Active": ("عدد الصفوف المفعّلة", '=COUNTIF(ClassActive,"نعم")', None),
        "Cls_Confirmed": ("صفوف محصورة لهذا اليوم", '=IF(Dash_DayIdx=0,0,' + '+'.join(f'(INDEX({cq(cn)}${DC0}$7:${DC1}$7,Dash_DayIdx)="✓")' for cn in CLASSES) + ')', None),
        "Best_Class": ("الصف الأعلى حضورًا", f'=IFERROR(INDEX({cq(DASH)}$F$31:$F$37,MATCH(MAX({cq(DASH)}$J$31:$J$37),{cq(DASH)}$J$31:$J$37,0)),"—")', None),
        "Low_Class": ("الصف الأقل حضورًا", f'=IFERROR(INDEX({cq(DASH)}$F$31:$F$37,MATCH(MIN({cq(DASH)}$J$31:$J$37),{cq(DASH)}$J$31:$J$37,0)),"—")', None),
    }
    r = 2
    for name, (lab, fm, nf) in cells.items():
        put(ws, f"A{r}", lab, f=font(9), al=ALIGN_R); put(ws, f"B{r}", fm, f=font(9, True), al=ALIGN_C, nf=nf)
        define(wb, name, f"{cq(CALC)}$B${r}"); r += 1
    # donut data
    r += 1
    put(ws, f"A{r}", "حضور", f=font(9)); put(ws, f"B{r}", '=IF(Dash_Rate="",0,Dash_Rate)', nf="0.0%"); d0 = r; r += 1
    put(ws, f"A{r}", "غياب", f=font(9)); put(ws, f"B{r}", '=IF(Dash_Rate="",0,1-Dash_Rate)', nf="0.0%"); r += 2
    # last 5 recorded days
    put(ws, f"A{r}", "آخر 5 أيام مسجَّلة", f=font(9, True, NAVY)); r += 1
    header_row(ws, r, 1, ["التاريخ", "التسمية", "الحاضرون", "المسجَّلون", "نسبة الحضور"], bg=MUTED, height=18); r += 1
    w0 = r
    for k in range(1, 6):
        n = f"(Rec_Days-5+{k})"
        put(ws, f"A{r}", f'=IF({n}<1,"",INDEX(Cal_Date,MATCH({n},Cal_RecIdx,0)))', nf="dd/mm/yyyy", f=font(9))
        put(ws, f"B{r}", f'=IF(A{r}="","",TEXT(A{r},"dd/mm")&" "&INDEX(Cal_Day,MATCH(A{r},Cal_Date,0)))', f=font(9))
        put(ws, f"C{r}", f'=IF(A{r}="","",IF(Dash_Class="الكل",INDEX(Cal_Pres,MATCH(A{r},Cal_Date,0))+INDEX(Cal_Late,MATCH(A{r},Cal_Date,0)),{class_counts("Dash_Class","P",f"A{r}")}+{class_counts("Dash_Class","L",f"A{r}")}))', f=font(9))
        put(ws, f"D{r}", f'=IF(A{r}="","",IF(Dash_Class="الكل",INDEX(Cal_Reg,MATCH(A{r},Cal_Date,0)),{class_counts("Dash_Class","P",f"A{r}")}+{class_counts("Dash_Class","A",f"A{r}")}+{class_counts("Dash_Class","E",f"A{r}")}+{class_counts("Dash_Class","L",f"A{r}")}))', f=font(9))
        put(ws, f"E{r}", f'=IF(OR(A{r}="",D{r}=0),"",IF(Dash_Class="الكل",INDEX(Cal_Rate,MATCH(A{r},Cal_Date,0)),({class_counts("Dash_Class","P",f"A{r}")}+IF(LateAsPresent="نعم",{class_counts("Dash_Class","L",f"A{r}")},0))/D{r}))', f=font(9), nf="0.0%")
        r += 1
    r += 1
    put(ws, f"A{r}", "مراجع القوائم (رقم الصف في جدول الطلاب)", f=font(9, True, NAVY)); r += 1
    header_row(ws, r, 1, ["الأعلى حضورًا", "يحتاجون دعمًا", "الحالات الحرجة", "تنبيهات اليوم"], bg=MUTED, height=18); r += 1
    i0 = r
    for k in range(1, 11):
        put(ws, f"A{r}", f'=IFERROR(MATCH(LARGE(Stu_KeyTop,{k}),Stu_KeyTop,0),"")', f=font(9))
        put(ws, f"B{r}", f'=IFERROR(MATCH(LARGE(Stu_KeySup,{k}),Stu_KeySup,0),"")', f=font(9))
        put(ws, f"C{r}", f'=IFERROR(MATCH(LARGE(Stu_KeyCrit,{k}),Stu_KeyCrit,0),"")', f=font(9))
        put(ws, f"D{r}", f'=IFERROR(MATCH({k},Stu_AlertSeq,0),"")', f=font(9))
        r += 1
    ws.sheet_state = "hidden"
    protect(ws)
    return {"donut": d0, "week": w0, "idx": i0}


def card(ws, col0, row0, label, value, sub, nf=None, big=20, vcolor=NAVY):
    for rr in range(row0, row0 + 3):
        for cc in (col0, col0 + 1):
            c = ws.cell(rr, cc); c.fill = fill(LIGHT)
            c.border = Border(left=side(WHITE, "medium") if cc == col0 + 1 else None, right=side(WHITE, "medium") if cc == col0 else None,
                              top=side(WHITE, "medium") if rr == row0 else None, bottom=side(WHITE, "medium") if rr == row0 + 2 else None)
    put(ws, (row0, col0), label, f=font(9, True, MUTED), al=Alignment(horizontal="centerContinuous", vertical="center", wrap_text=True))
    ws.cell(row0, col0 + 1).alignment = Alignment(horizontal="centerContinuous", vertical="center")
    put(ws, (row0 + 1, col0), value, f=font(big, True, vcolor), al=Alignment(horizontal="centerContinuous", vertical="center"), nf=nf)
    ws.cell(row0 + 1, col0 + 1).alignment = Alignment(horizontal="centerContinuous", vertical="center")
    put(ws, (row0 + 2, col0), sub, f=font(8, False, MUTED), al=Alignment(horizontal="centerContinuous", vertical="center"))
    ws.cell(row0 + 2, col0 + 1).alignment = Alignment(horizontal="centerContinuous", vertical="center")


def list_header(ws, row, col0, labels, bg=TEAL):
    for i, lab in enumerate(labels):
        put(ws, (row, col0 + i), lab, f=font(9, True, WHITE), bg=bg, al=ALIGN_CW, b=box(bg))
    ws.row_dimensions[row].height = 28


def build_dashboard(wb, calc_rows):
    ws = wb.create_sheet(DASH, 0)
    setup_sheet(ws, zoom=85, tab=NAVY)
    ws.column_dimensions["A"].width = 17
    for c in range(2, 17):
        ws.column_dimensions[L(c)].width = 11.5
    # title
    ws.row_dimensions[1].height = 30
    letterhead(ws, (7, 10), rows=(1, 3))
    put(ws, "A3", '="نظام متابعة وتحليل الحضور والغياب – قسم "&SectionName&"   |   Smart Attendance Analytics System"', f=font(14, True, NAVY), al=ALIGN_R)
    put(ws, "A4", '="العام الدراسي: "&AcademicYear&"   |   "&CurrentTerm&"   |   آخر يوم مسجَّل: "&TEXT(LastRecorded,"dd/mm/yyyy")&"   |   أيام مسجَّلة: "&Rec_Days', f=font(9, True, NAVY), al=ALIGN_R)
    ws.row_dimensions[2].height = 20; ws.row_dimensions[3].height = 24; ws.row_dimensions[4].height = 20
    # selectors
    put(ws, "A5", "التاريخ (فارغ = آخر يوم مسجَّل)", f=font(9, True, NAVY), al=ALIGN_R)
    input_cell(ws, "B5", None, nf="dd/mm/yyyy")
    put(ws, "D5", "الصف", f=font(9, True, NAVY), al=ALIGN_R)
    input_cell(ws, "E5", "الكل")
    put(ws, "G5", "الفترة (لنسبة الحضور العامة)", f=font(9, True, NAVY), al=ALIGN_R)
    input_cell(ws, "I5", "الشهر")
    put(ws, "K5", '="يعرض بيانات يوم: "&TEXT(Dash_Date,"dd/mm/yyyy")&IF(Dash_DayIdx=0,"  (⚠ ليس يومًا في التقويم)","  – "&INDEX(Cal_Day,Dash_DayIdx))&IF(Today_Reg=0,"  ⚠ لم يُسجَّل حضور هذا اليوم","")', f=font(10, True, TEAL), al=ALIGN_R)
    ws.row_dimensions[5].height = 22
    define(wb, "Dash_DateInput", f"{cq(DASH)}$B$5"); define(wb, "Dash_Class", f"{cq(DASH)}$E$5"); define(wb, "Dash_Period", f"{cq(DASH)}$I$5")
    dv = DataValidation(type="list", formula1="=FilterClassList", allow_blank=False, error="اختاري صفًا أو «الكل»", errorTitle="قيمة غير صحيحة"); ws.add_data_validation(dv); dv.add(ws["E5"])
    dv2 = DataValidation(type="list", formula1="=PeriodList", allow_blank=False); ws.add_data_validation(dv2); dv2.add(ws["I5"])
    dv3 = DataValidation(type="list", formula1="=Cal_Date", allow_blank=True, error="اختاري يوم دوام من القائمة أو اتركي الخلية فارغة", errorTitle="تاريخ غير صحيح", prompt="اختاري أي تاريخ سابق لعرض بياناته؛ فارغ = آخر يوم مسجَّل", promptTitle="التاريخ"); ws.add_data_validation(dv3); dv3.add(ws["B5"])
    # KPI cards
    cards = [("إجمالي أطفال KG1", "=Total_Kids", '="نشط – "&Dash_Class', None, NAVY),
             ("الحاضرون اليوم", "=Today_Att", '="منهم متأخرون: "&Today_L', None, GREEN_T),
             ("الغائبون اليوم", "=Today_A+Today_E", '="منهم بعذر: "&Today_E', None, ABS_T),
             ("نسبة الحضور اليوم", '=IF(Today_Rate="","—",Today_Rate)', '="من "&Today_Reg&" مسجَّل"', "0.0%", TEAL),
             ('="حالات "&Thr_FollowUp&"+ غياب"', '=COUNTIFS(Stu_InFilter,1,Stu_Level,">="&Thr_FollowUp)', '="جديدة اليوم: "&COUNTIF(Stu_LvlToday,Thr_FollowUp)', None, YEL_T),
             ('="حالات "&Thr_Repeated&"+ غياب"', '=COUNTIFS(Stu_InFilter,1,Stu_Level,">="&Thr_Repeated)', '="جديدة اليوم: "&COUNTIF(Stu_LvlToday,Thr_Repeated)', None, ORG_T),
             ('="حالات "&Thr_High&"+ غياب"', '=COUNTIFS(Stu_InFilter,1,Stu_Level,">="&Thr_High)', '="جديدة اليوم: "&COUNTIF(Stu_LvlToday,Thr_High)', None, DORG_T),
             ('="حالات "&Thr_Critical&"+ غياب"', '=COUNTIFS(Stu_InFilter,1,Stu_Level,">="&Thr_Critical)', '="جديدة اليوم: "&COUNTIF(Stu_LvlToday,Thr_Critical)', None, RED_T)]
    for j, (lab, val, sub, nf, colr) in enumerate(cards):
        card(ws, 1 + 2 * j, 7, lab, val, sub, nf=nf, vcolor=colr)
    ws.row_dimensions[7].height = 26; ws.row_dimensions[8].height = 34; ws.row_dimensions[9].height = 18
    ws.conditional_formatting.add("O7:P9", FormulaRule(formula=['$O$8>0'], fill=fill(RED_F)))
    ws.conditional_formatting.add("M7:N9", FormulaRule(formula=['$M$8>0'], fill=fill(DORG_F)))
    # ---- alerts box (A..G) & overall (I..P)
    section(ws, "A11", "تنبيهات المتابعة اليومية", "Today's Follow-Up Alerts", 7)
    section(ws, "I11", "نسبة الحضور العامة", "KG1 Overall Attendance", 8)
    lines = [("Thr_FollowUp", YEL_T), ("Thr_Repeated", ORG_T), ("Thr_High", DORG_T), ("Thr_Critical", RED_T)]
    for i, (thr, colr) in enumerate(lines):
        put(ws, f"A{12+i}", f'="• "&IF(COUNTIF(Stu_LvlToday,{thr})=0,"لا يوجد من وصل اليوم إلى ",IF(COUNTIF(Stu_LvlToday,{thr})=1,"طفل واحد وصل اليوم إلى ",COUNTIF(Stu_LvlToday,{thr})&" أطفال وصلوا اليوم إلى "))&{thr}&" أيام غياب"', f=font(10, True, colr), al=ALIGN_R)
    put(ws, "A16", '="الحالات الجديدة التي انتقلت اليوم إلى مستوى أعلى ("&TEXT(Dash_Date,"dd/mm/yyyy")&") – تواصلي مع أولياء أمورهم مباشرة:"', f=font(9, False, MUTED, True), al=ALIGN_R)
    list_header(ws, 17, 1, ["اسم الطفل", "الصف", "المستوى الذي بلغه", "إجمالي الغياب", "رقم ولي الأمر", "تم التواصل؟", "الإجراء"])
    i0 = calc_rows["idx"]
    for k in range(10):
        r = 18 + k; ref = f"{cq(CALC)}$D${i0+k}"
        G = lambda e: f'=IF({ref}="","",{e})'
        put(ws, f"A{r}", G(f"INDEX(Stu_Name,{ref})"), f=font(10, True), al=ALIGN_R, b=bottom())
        put(ws, f"B{r}", G(f"INDEX(Stu_Class,{ref})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"C{r}", G(f'"بلغ "&INDEX(Stu_LvlToday,{ref})&" غيابات"'), f=font(10, True, ORG_T), al=ALIGN_C, b=bottom())
        put(ws, f"D{r}", G(f"INDEX(Stu_Abs,{ref})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"E{r}", G(f"INDEX(Stu_Phone,{ref})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"F{r}", G(f"INDEX(Stu_Contacted,{ref})"), f=font(10, True), al=ALIGN_C, b=bottom())
        put(ws, f"G{r}", G(f"INDEX(Stu_Action,{ref})"), f=font(9), al=ALIGN_R, b=bottom())
        put(ws, f"H{r}", G(f"INDEX(Stu_LvlToday,{ref})"), f=font(8, False, MUTED), al=ALIGN_C)
    put(ws, "H17", "المستوى", f=font(8, True, WHITE), bg=MUTED, al=ALIGN_C)
    ws.conditional_formatting.add("A18:G27", FormulaRule(formula=['AND($A18<>"",$H18>=Thr_Critical)'], fill=fill(RED_F)))
    ws.conditional_formatting.add("A18:G27", FormulaRule(formula=['AND($A18<>"",$H18=Thr_High)'], fill=fill(DORG_F)))
    ws.conditional_formatting.add("F18:F27", FormulaRule(formula=['$F18="لا"'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add("F18:F27", FormulaRule(formula=['$F18="نعم"'], fill=fill(GREEN_F), font=Font(name=FONT, bold=True, color=GREEN_T)))
    put(ws, "A28", '=IF(COUNT(Stu_AlertSeq)=0,"لا توجد حالات جديدة اليوم ✓",IF(COUNT(Stu_AlertSeq)>10,"⚠ يوجد "&COUNT(Stu_AlertSeq)&" حالة جديدة – تُعرض أول 10؛ راجعي قائمة المتابعة",""))', f=font(9, True, GREEN_T), al=ALIGN_R)
    # overall
    put(ws, "I12", "KG1 Overall Attendance", f=font(9, False, MUTED, True), al=Alignment(horizontal="centerContinuous"))
    ws["J12"].alignment = Alignment(horizontal="centerContinuous"); ws["K12"].alignment = Alignment(horizontal="centerContinuous")
    put(ws, "I13", '=IF(Dash_Rate="","—",Dash_Rate)', f=font(30, True, TEAL), al=Alignment(horizontal="centerContinuous", vertical="center"), nf="0.0%")
    ws["J13"].alignment = Alignment(horizontal="centerContinuous", vertical="center"); ws["K13"].alignment = Alignment(horizontal="centerContinuous", vertical="center")
    ws.row_dimensions[13].height = 44
    put(ws, "I14", "=Dash_PeriodText", f=font(9, False, MUTED), al=Alignment(horizontal="centerContinuous"))
    ws["J14"].alignment = Alignment(horizontal="centerContinuous"); ws["K14"].alignment = Alignment(horizontal="centerContinuous")
    put(ws, "I15", '="الصف: "&Dash_Class', f=font(9, False, MUTED), al=Alignment(horizontal="centerContinuous"))
    ws["J15"].alignment = Alignment(horizontal="centerContinuous"); ws["K15"].alignment = Alignment(horizontal="centerContinuous")
    put(ws, "I17", "توزيع الأطفال حسب فئة الحضور (العام كاملًا)", f=font(9, True, NAVY), al=ALIGN_R)
    list_header(ws, 18, 9, ["الفئة", "عدد الأطفال", "النسبة"])
    for i, (lab, colr) in enumerate([("Cat_L1", GREEN_T), ("Cat_L2", GREEN_T), ("Cat_L3", YEL_T), ("Cat_L4", RED_T)]):
        r = 19 + i
        put(ws, f"I{r}", f"={lab}", f=font(10, True, colr), al=ALIGN_R, b=bottom())
        put(ws, f"J{r}", f"=COUNTIFS(Stu_InFilter,1,Stu_Cat,{lab})", f=font(11, True, NAVY), al=ALIGN_C, b=bottom())
        put(ws, f"K{r}", f'=IF(Total_Kids=0,"",J{r}/Total_Kids)', f=font(10), al=ALIGN_C, b=bottom(), nf="0%")
    put(ws, "I23", "«حضور كامل» = بدون أي غياب. الفئات الأخرى حسب نسبة الحضور التراكمية للعام.", f=font(8, False, MUTED, True), al=ALIGN_R)
    # donut chart
    calc = wb[CALC]
    ch = DoughnutChart(holeSize=62)
    d0 = calc_rows["donut"]
    data = Reference(calc, min_col=2, min_row=d0, max_row=d0 + 1)
    cats = Reference(calc, min_col=1, min_row=d0, max_row=d0 + 1)
    ch.add_data(data, titles_from_data=False); ch.set_categories(cats)
    s = ch.series[0]
    for idx, colr in [(0, TEAL), (1, "E5E7EB")]:
        pt = DataPoint(idx=idx); pt.graphicalProperties.solidFill = colr; pt.graphicalProperties.line.solidFill = WHITE; s.dPt.append(pt)
    ch.legend = None; ch.title = None
    place(ws, ch, 13, 12, 16, 22)
    # ---- today's analysis & class comparison
    section(ws, "A29", "تحليل اليوم", "Today's Attendance Analysis", 4)
    section(ws, "F29", "مقارنة الصفوف", "Attendance by Class", 8)
    rows = [("إجمالي الأطفال", "=Total_Kids", None), ("الحاضرون (حاضر + متأخر)", "=Today_Att", None), ("الغائبون (بدون عذر)", "=Today_A", None),
            ("الغياب بعذر", "=Today_E", None), ("المتأخرون", "=Today_L", None), ("نسبة الحضور اليوم", '=IF(Today_Rate="","—",Today_Rate)', "0.0%"),
            ("الصف الأعلى حضورًا", "=Best_Class", None), ("الصف الأقل حضورًا", "=Low_Class", None),
            ("الصفوف المحصورة لهذا اليوم", '=Cls_Confirmed&" من "&Cls_Active', None)]
    for i, (lab, fm, nf) in enumerate(rows):
        r = 30 + i
        put(ws, f"A{r}", lab, f=font(10), al=ALIGN_R, b=bottom())
        put(ws, f"B{r}", fm, f=font(11, True, NAVY), al=ALIGN_C, b=bottom(), nf=nf)
    ws.conditional_formatting.add("B38", FormulaRule(formula=['Cls_Confirmed<Cls_Active'], font=Font(name=FONT, bold=True, color=ORG_T)))
    list_header(ws, 30, 6, ["الصف", "عدد الأطفال", "الحاضرون اليوم", "الغائبون اليوم", "نسبة الحضور اليوم", "نسبة الحضور الشهرية", '="حالات "&Thr_FollowUp&"+"', '="غياب مرتفع "&Thr_High&"+"'])
    for k in range(NCLS):
        r = 31 + k; kk = k + 1
        put(ws, f"F{r}", f'=IF(INDEX(ClassActive,{kk})="نعم",INDEX(ClassList,{kk}),"")', f=font(10, True), al=ALIGN_C, b=bottom())
        G = lambda e: f'=IF($F{r}="","",{e})'
        put(ws, f"G{r}", G(f'COUNTIFS(Stu_Class,$F{r},Stu_StStatus,"نشط")'), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"H{r}", G(f'COUNTIFS(Stu_Class,$F{r},Stu_Today,"P",Stu_StStatus,"نشط")+COUNTIFS(Stu_Class,$F{r},Stu_Today,"L",Stu_StStatus,"نشط")'), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"I{r}", G(f'COUNTIFS(Stu_Class,$F{r},Stu_Today,"A",Stu_StStatus,"نشط")+COUNTIFS(Stu_Class,$F{r},Stu_Today,"E",Stu_StStatus,"نشط")'), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"J{r}", G(f'IF(H{r}+I{r}=0,"",(COUNTIFS(Stu_Class,$F{r},Stu_Today,"P",Stu_StStatus,"نشط")+IF(LateAsPresent="نعم",COUNTIFS(Stu_Class,$F{r},Stu_Today,"L",Stu_StStatus,"نشط"),0))/(H{r}+I{r}))'), f=font(10, True, TEAL), al=ALIGN_C, b=bottom(), nf="0.0%")
        att = f'COUNTIFS(DB_Class,$F{r},DB_Date,">="&Dash_MStart,DB_Date,"<="&Dash_MEnd,DB_Code,"P")+IF(LateAsPresent="نعم",COUNTIFS(DB_Class,$F{r},DB_Date,">="&Dash_MStart,DB_Date,"<="&Dash_MEnd,DB_Code,"L"),0)'
        reg = "+".join(f'COUNTIFS(DB_Class,$F{r},DB_Date,">="&Dash_MStart,DB_Date,"<="&Dash_MEnd,DB_Code,"{c}")' for c in "PAEL")
        put(ws, f"K{r}", G(f'IF(({reg})=0,"",({att})/({reg}))'), f=font(10), al=ALIGN_C, b=bottom(), nf="0.0%")
        put(ws, f"L{r}", G(f'COUNTIFS(Stu_Class,$F{r},Stu_StStatus,"نشط",Stu_Level,">="&Thr_FollowUp)'), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"M{r}", G(f'COUNTIFS(Stu_Class,$F{r},Stu_StStatus,"نشط",Stu_Level,">="&Thr_High)'), f=font(10), al=ALIGN_C, b=bottom())
    ws.conditional_formatting.add("M31:M37", FormulaRule(formula=['AND($M31<>"",$M31>0)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add("L31:L37", FormulaRule(formula=['AND($L31<>"",$L31>0)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add("J31:J37", FormulaRule(formula=['AND($J31<>"",$J31<Cat_90)'], font=Font(name=FONT, bold=True, color=RED_T)))
    put(ws, "F38", "نسبة الحضور الشهرية = شهر التاريخ المحدد. الصفوف الفارغة = أوراق احتياطية غير مفعّلة.", f=font(8, False, MUTED, True), al=ALIGN_R)
    # bar chart by class
    bc = BarChart(); bc.type = "col"; bc.style = 10
    bc.title = "Attendance Rate by Class – نسبة الحضور اليوم حسب الصف"
    data = Reference(ws, min_col=10, min_row=30, max_row=37); cats = Reference(ws, min_col=6, min_row=31, max_row=37)
    bc.add_data(data, titles_from_data=True); bc.set_categories(cats)
    bc.series[0].graphicalProperties.solidFill = TEAL
    bc.y_axis.scaling.min = 0; bc.y_axis.scaling.max = 1; bc.y_axis.number_format = "0%"; bc.y_axis.majorGridlines = None
    labels(bc)
    bc.legend = None
    place(ws, bc, 6, 40, 13, 52)
    # last 5 recorded days: visible mini table + line chart
    list_header(ws, 40, 1, ["التاريخ", "اليوم", "الحاضرون", "المسجَّلون", "نسبة الحضور"])
    for k in range(1, 6):
        r = 40 + k; n = f"(Rec_Days-5+{k})"
        put(ws, f"A{r}", f'=IF({n}<1,"",INDEX(Cal_Date,MATCH({n},Cal_RecIdx,0)))', f=font(9, True), al=ALIGN_C, b=bottom(), nf="dd/mm/yyyy")
        put(ws, f"B{r}", f'=IF(A{r}="","",INDEX(Cal_Day,MATCH(A{r},Cal_Date,0)))', f=font(9), al=ALIGN_C, b=bottom())
        put(ws, f"C{r}", f'=IF(A{r}="","",IF(Dash_Class="الكل",INDEX(Cal_Pres,MATCH(A{r},Cal_Date,0))+INDEX(Cal_Late,MATCH(A{r},Cal_Date,0)),{class_counts("Dash_Class","P",f"A{r}")}+{class_counts("Dash_Class","L",f"A{r}")}))', f=font(9), al=ALIGN_C, b=bottom())
        put(ws, f"D{r}", f'=IF(A{r}="","",IF(Dash_Class="الكل",INDEX(Cal_Reg,MATCH(A{r},Cal_Date,0)),{class_counts("Dash_Class","P",f"A{r}")}+{class_counts("Dash_Class","A",f"A{r}")}+{class_counts("Dash_Class","E",f"A{r}")}+{class_counts("Dash_Class","L",f"A{r}")}))', f=font(9), al=ALIGN_C, b=bottom())
        put(ws, f"E{r}", f'=IF(OR(A{r}="",D{r}=0),"",IF(Dash_Class="الكل",INDEX(Cal_Rate,MATCH(A{r},Cal_Date,0)),({class_counts("Dash_Class","P",f"A{r}")}+IF(LateAsPresent="نعم",{class_counts("Dash_Class","L",f"A{r}")},0))/D{r}))', f=font(9, True, TEAL), al=ALIGN_C, b=bottom(), nf="0.0%")
    lc = LineChart(); lc.title = "Attendance Trend – آخر 5 أيام مسجَّلة"; lc.style = 12
    data = Reference(ws, min_col=5, min_row=41, max_row=45); cats = Reference(ws, min_col=1, min_row=41, max_row=45)
    lc.add_data(data, titles_from_data=False); lc.set_categories(cats)
    s = lc.series[0]; s.graphicalProperties.line.solidFill = TEAL; s.graphicalProperties.line.width = 28000; s.marker.symbol = "circle"; s.marker.size = 7
    s.marker.graphicalProperties.solidFill = TEAL; s.smooth = False
    lc.y_axis.scaling.min = 0; lc.y_axis.scaling.max = 1; lc.y_axis.number_format = "0%"; lc.x_axis.number_format = "dd/mm"
    labels(lc)
    lc.legend = None
    place(ws, lc, 2, 47, 5, 58)
    # ---- lists
    section(ws, "A60", "الأعلى حضورًا ⭐", "Top Attendance", 5)
    section(ws, "F60", "أطفال يحتاجون دعمًا في الحضور", "Students Requiring Attendance Support", 6)
    section(ws, "L60", "حالات الغياب الحرجة 🔴", "Critical Absence Cases", 5)
    list_header(ws, 61, 1, ["اسم الطفل", "الصف", "أيام الحضور", "نسبة الحضور", "الفئة"])
    list_header(ws, 61, 6, ["اسم الطفل", "الصف", "أيام الغياب", "نسبة الحضور", "حالة المتابعة", "تم التواصل؟"], bg=ORG_T)
    list_header(ws, 61, 12, ["اسم الطفل", "الصف", "الغياب", "رقم ولي الأمر", "تم التواصل؟"], bg=RED_T)
    for k in range(10):
        r = 62 + k
        a = f"{cq(CALC)}$A${i0+k}"; b = f"{cq(CALC)}$B${i0+k}"; c = f"{cq(CALC)}$C${i0+k}"
        GA = lambda e: f'=IF({a}="","",{e})'; GB = lambda e: f'=IF({b}="","",{e})'; GC = lambda e: f'=IF({c}="","",{e})'
        put(ws, f"A{r}", GA(f'IF({k}<3,"⭐ ","")&INDEX(Stu_Name,{a})'), f=font(10, True), al=ALIGN_R, b=bottom())
        put(ws, f"B{r}", GA(f"INDEX(Stu_Class,{a})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"C{r}", GA(f"INDEX(Stu_Pres,{a})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"D{r}", GA(f"INDEX(Stu_Rate,{a})"), f=font(10, True, GREEN_T), al=ALIGN_C, b=bottom(), nf="0.0%")
        put(ws, f"E{r}", GA(f"INDEX(Stu_Cat,{a})"), f=font(9), al=ALIGN_C, b=bottom())
        put(ws, f"F{r}", GB(f"INDEX(Stu_Name,{b})"), f=font(10, True), al=ALIGN_R, b=bottom())
        put(ws, f"G{r}", GB(f"INDEX(Stu_Class,{b})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"H{r}", GB(f"INDEX(Stu_Abs,{b})+INDEX(Stu_Exc,{b})"), f=font(10, True, ABS_T), al=ALIGN_C, b=bottom())
        put(ws, f"I{r}", GB(f"INDEX(Stu_Rate,{b})"), f=font(10), al=ALIGN_C, b=bottom(), nf="0.0%")
        put(ws, f"J{r}", GB(f"INDEX(Stu_Status,{b})"), f=font(9, True), al=ALIGN_C, b=bottom())
        put(ws, f"K{r}", GB(f"INDEX(Stu_Contacted,{b})"), f=font(10, True), al=ALIGN_C, b=bottom())
        put(ws, f"L{r}", GC(f"INDEX(Stu_Name,{c})"), f=font(10, True, RED_T), al=ALIGN_R, b=bottom())
        put(ws, f"M{r}", GC(f"INDEX(Stu_Class,{c})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"N{r}", GC(f"INDEX(Stu_Abs,{c})"), f=font(10, True, RED_T), al=ALIGN_C, b=bottom())
        put(ws, f"O{r}", GC(f"INDEX(Stu_Phone,{c})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"P{r}", GC(f"INDEX(Stu_Contacted,{c})"), f=font(10, True), al=ALIGN_C, b=bottom())
    for rg in ["K62:K71", "P62:P71"]:
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{rg[0]}62="لا"'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{rg[0]}62="نعم"'], fill=fill(GREEN_F), font=Font(name=FONT, bold=True, color=GREEN_T)))
    ws.conditional_formatting.add("J62:J71", FormulaRule(formula=['ISNUMBER(SEARCH("حرجة",$J62))'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.conditional_formatting.add("J62:J71", FormulaRule(formula=['ISNUMBER(SEARCH("مرتفع",$J62))'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add("J62:J71", FormulaRule(formula=['ISNUMBER(SEARCH("متكرر",$J62))'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add("J62:J71", FormulaRule(formula=['ISNUMBER(SEARCH("متابعة",$J62))'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add("L62:P71", FormulaRule(formula=['$L62<>""'], fill=fill("FDECEC")))
    put(ws, "L72", '=IF(COUNT(Stu_KeyCrit)=0,"لا توجد حالات حرجة ✓","")', f=font(9, True, GREEN_T), al=ALIGN_R)
    # links
    put(ws, "A74", "انتقال سريع:", f=font(9, True, NAVY), al=ALIGN_R)
    links = [("B74", SHEETS["daily"], "الحضور اليومي"), ("D74", SHEETS["fu"], "قائمة المتابعة"), ("F74", SHEETS["alog"], "سجل الغياب"), ("H74", SHEETS["clog"], "سجل التواصل"),
             ("J74", SHEETS["prof"], "ملف الطفل"), ("L74", SHEETS["trend"], "الاتجاهات"), ("N74", SHEETS["rpt"], "🖨 طباعة تقرير يومي/أسبوعي/شهري"), ("P74", SHEETS["guide"], "الدليل")]
    for ref, sh, lab in links:
        put(ws, ref, f'=HYPERLINK("#{q(sh)}!A1","{lab} ⬅")', f=font(10, True, "1D4ED8"), al=ALIGN_R)
    approval_block(ws, "A76", "A77", "A78")
    principal_block(ws, "O76", "O77", "O78")
    ws.freeze_panes = "A6"
    ws.print_area = "A1:P78"
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    protect(ws)
    return ws
