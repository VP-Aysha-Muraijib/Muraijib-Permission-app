# -*- coding: utf-8 -*-
import sys, time
from openpyxl import Workbook
from common import *
from settings_cal import build_settings, build_calendar
from registers import build_register
from master_db import build_master, build_db, build_contact_log, build_absence_log, build_followup
from dashboard import build_calc, build_dashboard
from others import build_daily, build_profile, build_trends
from report_guide import build_report, build_guide
import testdata

def build(path, with_demo=True):
    t0 = time.time()
    wb = Workbook(); wb.remove(wb.active)
    build_settings(wb)
    build_calendar(wb)
    regs = [build_register(wb, cn, i) for i, cn in enumerate(CLASSES)]
    build_master(wb)
    build_db(wb)
    build_contact_log(wb)
    build_absence_log(wb)
    build_followup(wb)
    calc_rows = build_calc(wb)
    build_dashboard(wb, calc_rows)
    build_daily(wb)
    build_profile(wb)
    build_trends(wb)
    build_report(wb)
    build_guide(wb)
    # sheet order
    order = [SHEETS["dash"], SHEETS["daily"]] + CLASSES + [SHEETS["fu"], SHEETS["alog"], SHEETS["clog"], SHEETS["prof"], SHEETS["trend"], SHEETS["rpt"], SHEETS["stu"], SHEETS["cal"], SHEETS["set"], SHEETS["guide"], SHEETS["db"], SHEETS["calc"]]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    if with_demo:
        students = testdata.make_students()
        for ci, cn in enumerate(CLASSES[:ACTIVE_CLASSES]):
            ws = wb[cn]
            for i, s in enumerate(students[cn]):
                r = REG_FIRST + i
                ws[f"B{r}"] = s["id"]; ws[f"C{r}"] = s["name"]; ws[f"I{r}"] = s["status"]; ws[f"K{r}"] = s["guardian"]; ws[f"L{r}"] = s["phone"]
                if s["join"]: ws[f"J{r}"] = s["join"]
                for d, v in enumerate(s["grid"]):
                    if v: ws.cell(r, DATE_COL0 + d, v)
            for d in range(NDAYS):
                if testdata.confirmed(ci, d): ws.cell(7, DATE_COL0 + d, "✓")
            ws[f"M{REG_FIRST+3}"] = "متابعة مستمرة مع الأسرة"
        logs = testdata.make_contacts(students)
        wl = wb[SHEETS["clog"]]
        for k, (d, sid, m, note) in enumerate(logs):
            r = LOG_FIRST + k
            wl[f"A{r}"] = d; wl[f"B{r}"] = sid; wl[f"F{r}"] = m; wl[f"G{r}"] = note
        wb[SHEETS["prof"]]["B4"] = students[CLASSES[0]][5]["name"]
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    print("saved", path, "in", round(time.time() - t0, 1), "s")
    return students if with_demo else None

if __name__ == "__main__":
    build(sys.argv[1], with_demo=True)
    if len(sys.argv) > 2:
        build(sys.argv[2], with_demo=False)
