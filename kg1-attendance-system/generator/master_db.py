# -*- coding: utf-8 -*-
from common import *
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

STU_HDR = ["م", "الرقم الطلابي", "اسم الطفل", "الصف", "الشعبة", "الحالة الدراسية", "تاريخ الالتحاق", "اسم ولي الأمر", "رقم التواصل",
           "أيام الدوام", "الحضور", "الغياب", "غياب بعذر", "تأخير", "نسبة الحضور", "نسبة الغياب", "مستوى التنبيه", "حالة المتابعة", "آخر تاريخ غياب",
           "بلوغ حد المتابعة", "بلوغ الغياب المتكرر", "بلوغ الغياب المرتفع", "بلوغ الحالة الحرجة", "تاريخ آخر تواصل", "طريقة آخر تواصل", "تم التواصل؟",
           "التواصل مع ولي الأمر", "الإجراء المطلوب", "فئة الحضور", "ملاحظات", "رقم مكرر؟",
           "مشمول في فلتر اللوحة", "حالة اليوم المحدد", "حضور الفترة", "أيام الفترة", "مستوى بلغه اليوم", "تسلسل تنبيه اليوم",
           "مفتاح المتابعة", "مفتاح الأعلى حضورًا", "مفتاح الدعم", "مفتاح الحرجة", "مشمول في التقرير", "مفتاح التقرير", "حضور فترة التقرير", "أيام فترة التقرير", "غياب فترة التقرير", "رقم الصف", "رقم الصف الداخلي"]
STU_NAMES = {"B": "Stu_ID", "C": "Stu_Name", "D": "Stu_Class", "E": "Stu_Section", "F": "Stu_StStatus", "G": "Stu_Join", "H": "Stu_Guardian", "I": "Stu_Phone",
             "J": "Stu_Days", "K": "Stu_Pres", "L": "Stu_Abs", "M": "Stu_Exc", "N": "Stu_Late", "O": "Stu_Rate", "P": "Stu_AbsRate", "Q": "Stu_Level", "R": "Stu_Status",
             "S": "Stu_LastAbs", "T": "Stu_D3", "U": "Stu_D5", "V": "Stu_D10", "W": "Stu_D15", "X": "Stu_LastContact", "Y": "Stu_Method", "Z": "Stu_Contacted",
             "AA": "Stu_ContactReq", "AB": "Stu_Action", "AC": "Stu_Cat", "AD": "Stu_Notes", "AE": "Stu_Dup", "AF": "Stu_InFilter", "AG": "Stu_Today",
             "AH": "Stu_PAtt", "AI": "Stu_PTot", "AJ": "Stu_LvlToday", "AK": "Stu_AlertSeq", "AL": "Stu_KeyFU", "AM": "Stu_KeyTop", "AN": "Stu_KeySup",
             "AO": "Stu_KeyCrit", "AP": "Stu_RptIn", "AQ": "Stu_KeyRpt", "AR": "Stu_RAtt", "AS": "Stu_RTot", "AT": "Stu_RAbs", "AU": "Stu_C", "AV": "Stu_I"}

def period_counts(S, rr, start, end):
    """Return (attended_formula_part, total_formula_part, absent_part) for register row rr within [start,end]."""
    RNG = f"{cq(S)}${DC0}{rr}:${DC1}{rr}"
    D = f"{cq(S)}${DC0}$5:${DC1}$5"; SCH = f"{cq(S)}${DC0}$6:${DC1}$6"; CONF = f"{cq(S)}${DC0}$7:${DC1}$7"
    crit = f'{SCH},1,{D},">="&MAX({cq(S)}$W{rr},{start}),{D},"<="&{end}'
    P = f'COUNTIFS({RNG},"حاضر",{crit})+COUNTIFS({CONF},"✓",{crit})-COUNTIFS({RNG},"<>",{CONF},"✓",{crit})'
    A = f'COUNTIFS({RNG},"غائب",{crit})'; E = f'COUNTIFS({RNG},"بعذر",{crit})'; Lt = f'COUNTIFS({RNG},"متأخر",{crit})'
    att = f'{P}+IF(LateAsPresent="نعم",{Lt},0)'
    tot = f'{P}+{A}+{E}+{Lt}'
    ab = f'{A}+{E}'
    return att, tot, ab

def build_master(wb):
    ws = wb.create_sheet(SHEETS["stu"])
    setup_sheet(ws, tab="9CA3AF")
    title_block(ws, "قاعدة بيانات الأطفال الموحدة", "تُجمَّع تلقائيًا من أوراق الصفوف (tblStudents) — الإدخال يتم في ورقة كل صف، وهذه الورقة للعرض والفلترة والتحليل فقط", "Master student table (auto)")
    widths = [4, 11, 22, 9, 7, 9, 11, 16, 13, 6, 6, 6, 6, 6, 8, 8, 7, 14, 11, 11, 11, 11, 11, 11, 12, 8, 16, 30, 12, 20, 8] + [9] * 17
    header_row(ws, 4, 1, STU_HDR, widths=widths, height=42)
    for c in range(32, 49):
        ws.cell(4, c).fill = fill(MUTED); ws.cell(4, c).font = font(8, True, WHITE)
    for ci, S in enumerate(CLASSES):
        for i in range(1, SLOTS + 1):
            r = STU_FIRST + ci * SLOTS + i - 1
            rr = REG_FIRST + i - 1
            M = lambda col: f'=IF({cq(S)}${col}{rr}="","",{cq(S)}${col}{rr})'
            f = {"A": f'=IF(B{r}="","",ROW()-{STU_FIRST-1})', "B": M("B"), "C": M("C"), "D": f'=IF(B{r}="","",{cq(S)}$B$2)',
                 "E": f'=IF(B{r}="","",MID(D{r},FIND("-",D{r}&"-")+1,10))', "F": M("I"), "G": M("J"), "H": M("K"), "I": M("L"),
                 "J": M("N"), "K": M("D"), "L": M("E"), "M": M("O"), "N": M("P"), "O": M("F"), "P": M("Q"), "Q": M("V"), "R": M("G"), "S": M("R"),
                 "T": M("X"), "U": M("Y"), "V": M("Z"), "W": M("AA"), "X": M("T"), "Y": M("U"), "Z": M("S"), "AA": M("H"),
                 "AB": f'=IF(Q{r}="","",INDEX(LevelActions,MATCH(Q{r},LevelValues,0)))',
                 "AC": f'=IF(J{r}="","",IF(J{r}=0,"—",IF(AND(L{r}=0,M{r}=0),Cat_L1,IF(O{r}>=Cat_95,Cat_L2,IF(O{r}>=Cat_90,Cat_L3,Cat_L4)))))',
                 "AD": M("M"), "AE": f'=IF(B{r}="","",IF(COUNTIF($B${STU_FIRST}:$B${STU_LAST},B{r})>1,"مكرر!",""))',
                 "AF": f'=IF(AND(B{r}<>"",F{r}="نشط",OR(Dash_Class="الكل",D{r}=Dash_Class)),1,0)',
                 "AG": f'=IF(OR(AF{r}=0,Dash_DayIdx=0),"",INDEX(DB_Code,{1 + ci * DB_BLOCK + i - 1}+(Dash_DayIdx-1)*{SLOTS}))',
                 "AJ": f'=IF(AF{r}=0,0,IF(W{r}=Dash_Date,Thr_Critical,IF(V{r}=Dash_Date,Thr_High,IF(U{r}=Dash_Date,Thr_Repeated,IF(T{r}=Dash_Date,Thr_FollowUp,0)))))',
                 "AK": f'=IF(AJ{r}>0,COUNTIF($AJ${STU_FIRST}:AJ{r},">0"),"")',
                 "AL": f'=IF(AND(B{r}<>"",F{r}="نشط",Q{r}<>"",Q{r}>0),L{r}*1000+(999-ROW()),"")',
                 "AM": f'=IF(AND(AF{r}=1,J{r}<>"",J{r}>0),ROUND(O{r}*10000,0)*1000000+K{r}*1000+(999-ROW()),"")',
                 "AN": f'=IF(AND(AF{r}=1,J{r}<>"",J{r}>0),(L{r}+M{r})*10000000+ROUND((1-O{r})*10000,0)*1000+(999-ROW()),"")',
                 "AO": f'=IF(AND(AF{r}=1,Q{r}<>"",Q{r}>=Thr_Critical),L{r}*1000+(999-ROW()),"")',
                 "AP": f'=IF(AND(B{r}<>"",F{r}="نشط",OR(Rpt_Class="الكل",D{r}=Rpt_Class)),1,0)',
                 "AQ": f'=IF(AND(AP{r}=1,Q{r}<>"",Q{r}>0),L{r}*1000+(999-ROW()),"")',
                 "AU": ci + 1, "AV": i}
            att, tot, ab = period_counts(S, rr, "Dash_PStart", "Dash_PEnd")
            f["AH"] = f'=IF(AF{r}=0,"",{att})'; f["AI"] = f'=IF(AF{r}=0,"",{tot})'
            att2, tot2, ab2 = period_counts(S, rr, "Rpt_PStart", "Rpt_PEnd")
            f["AR"] = f'=IF(AP{r}=0,"",{att2})'; f["AS"] = f'=IF(AP{r}=0,"",{tot2})'; f["AT"] = f'=IF(AP{r}=0,"",{ab2})'
            for col, fm in f.items():
                c = put(ws, f"{col}{r}", fm, f=font(9), al=ALIGN_C)
            for col in ["C", "H", "R", "AB", "AD"]:
                ws[f"{col}{r}"].alignment = ALIGN_R
            for col in ["G", "S", "T", "U", "V", "W", "X"]:
                ws[f"{col}{r}"].number_format = "dd/mm/yyyy"
            ws[f"O{r}"].number_format = "0.0%"; ws[f"P{r}"].number_format = "0.0%"
            for col in list(STU_NAMES)[31:]:
                ws[f"{col}{r}"].font = font(8, False, MUTED)
    for col, name in STU_NAMES.items():
        define(wb, name, f"{cq(ws.title)}${col}${STU_FIRST}:${col}${STU_LAST}")
    tab = Table(displayName="tblStudents", ref=f"A4:AV{STU_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "D5"
    rg = f"A{STU_FIRST}:AD{STU_LAST}"
    ws.conditional_formatting.add(f"R{STU_FIRST}:R{STU_LAST}", FormulaRule(formula=[f'AND($Q{STU_FIRST}<>"",$Q{STU_FIRST}>=Thr_Critical)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.conditional_formatting.add(f"R{STU_FIRST}:R{STU_LAST}", FormulaRule(formula=[f'AND($Q{STU_FIRST}<>"",$Q{STU_FIRST}=Thr_High)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add(f"R{STU_FIRST}:R{STU_LAST}", FormulaRule(formula=[f'AND($Q{STU_FIRST}<>"",$Q{STU_FIRST}=Thr_Repeated)'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add(f"R{STU_FIRST}:R{STU_LAST}", FormulaRule(formula=[f'AND($Q{STU_FIRST}<>"",$Q{STU_FIRST}=Thr_FollowUp)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add(f"AE{STU_FIRST}:AE{STU_LAST}", FormulaRule(formula=[f'$AE{STU_FIRST}="مكرر!"'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.conditional_formatting.add(f"AA{STU_FIRST}:AA{STU_LAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("مطلوب",$AA{STU_FIRST}))'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    protect(ws)
    return ws


def build_db(wb):
    ws = wb.create_sheet(SHEETS["db"])
    setup_sheet(ws, tab="9CA3AF", grid=True)
    hdr = ["التاريخ", "الصف", "الرقم الطلابي", "اسم الطفل", "الحالة", "الرمز", "غياب بدون عذر", "الغياب التراكمي", "حدث غياب", "تسلسل الغياب"]
    header_row(ws, 1, 1, hdr, widths=[12, 9, 12, 22, 9, 6, 8, 9, 8, 9], height=30)
    r = DB_FIRST
    f10 = font(9)
    for ci, S in enumerate(CLASSES):
        P = cq(S)
        for di, d in enumerate(DATES):
            col = L(DATE_COL0 + di)
            for i in range(1, SLOTS + 1):
                rr = REG_FIRST + i - 1
                ws.cell(r, 1, d).number_format = "dd/mm/yyyy"
                ws.cell(r, 2, f"={P}$B$2")
                ws.cell(r, 3, f'=IF({P}$B{rr}="","",{P}$B{rr})')
                ws.cell(r, 4, f'=IF(C{r}="","",{P}$C{rr})')
                ws.cell(r, 5, f'=IF(OR(C{r}="",{P}$I{rr}<>"نشط",{P}{col}$6<>1,{P}$W{rr}>A{r}),"",IF({P}{col}{rr}<>"",{P}{col}{rr},IF({P}{col}$7="✓","حاضر","")))')
                ws.cell(r, 6, f'=IF(E{r}="","",IF(E{r}="حاضر","P",IF(E{r}="غائب","A",IF(E{r}="بعذر","E",IF(E{r}="متأخر","L","")))))')
                ws.cell(r, 7, f'=IF(F{r}="A",1,0)')
                ws.cell(r, 9, f'=IF(OR(F{r}="A",F{r}="E"),1,0)')
                ws.cell(r, 8, f'=IF(I{r}=1,COUNTIFS({P}${DC0}{rr}:{col}{rr},"غائب",{P}${DC0}$6:{col}$6,1,{P}${DC0}$5:{col}$5,">="&{P}$W{rr}),"")')
                ws.cell(r, 10, f'=I{r}' if r == DB_FIRST else f'=J{r-1}+I{r}')
                r += 1
    assert r - 1 == DB_LAST
    for col, name in {"A": "DB_Date", "B": "DB_Class", "C": "DB_ID", "D": "DB_Name", "E": "DB_Status", "F": "DB_Code", "G": "DB_IsAbs", "H": "DB_Cum", "I": "DB_IsEvt", "J": "DB_Seq"}.items():
        define(wb, name, f"{cq(ws.title)}${col}${DB_FIRST}:${col}${DB_LAST}")
    tab = Table(displayName="tblAttendance", ref=f"A1:J{DB_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    protect(ws)
    return ws


def build_contact_log(wb):
    ws = wb.create_sheet(SHEETS["clog"])
    setup_sheet(ws, tab="D97706")
    title_block(ws, "سجل التواصل مع أولياء الأمور", "سجّلي هنا كل تواصل مع ولي الأمر: التاريخ + الرقم الطلابي + الطريقة + النتيجة. يُحدَّث عمود «تم التواصل؟» في جميع الأوراق تلقائيًا", "Parent contact log (input)")
    hdr = ["التاريخ", "الرقم الطلابي", "اسم الطفل", "الصف", "عدد الغيابات عند التواصل", "طريقة التواصل", "نتيجة التواصل / ملاحظات", "المستوى عند التواصل", "مفتاح", "تسلسل الطفل المحدد"]
    header_row(ws, 5, 1, hdr, widths=[12, 12, 22, 9, 11, 14, 40, 14, 3, 3], height=36)
    for r in range(LOG_FIRST, LOG_LAST + 1):
        input_cell(ws, f"A{r}", None, nf="dd/mm/yyyy"); input_cell(ws, f"B{r}", None); input_cell(ws, f"F{r}", None); input_cell(ws, f"G{r}", None, al=ALIGN_R)
        put(ws, f"C{r}", f'=IF(B{r}="","",IFERROR(INDEX(Stu_Name,MATCH(B{r},Stu_ID,0)),"رقم غير موجود"))', f=font(10), al=ALIGN_R, b=box())
        put(ws, f"D{r}", f'=IF(B{r}="","",IFERROR(INDEX(Stu_Class,MATCH(B{r},Stu_ID,0)),""))', f=font(10), al=ALIGN_C, b=box())
        put(ws, f"E{r}", f'=IF(OR(B{r}="",A{r}=""),"",COUNTIFS(DB_ID,B{r},DB_Date,"<="&A{r},DB_IsAbs,1))', f=font(10), al=ALIGN_C, b=box())
        put(ws, f"H{r}", f'=IF(E{r}="","",INDEX(LevelLabels,MATCH(E{r},LevelValues,1)))', f=font(10), al=ALIGN_C, b=box())
        put(ws, f"I{r}", f'=IF(OR(B{r}="",A{r}=""),"",B{r}&"|"&A{r})', f=font(8, False, MUTED), al=ALIGN_C)
        put(ws, f"J{r}", f'=IF(AND(B{r}<>"",B{r}=Prof_ID),COUNTIF($B${LOG_FIRST}:B{r},Prof_ID),"")', f=font(8, False, MUTED), al=ALIGN_C)
    for col, name in {"A": "Log_Date", "B": "Log_ID", "C": "Log_Name", "E": "Log_Count", "F": "Log_Method", "G": "Log_Notes", "H": "Log_Level", "I": "Log_Key", "J": "Log_Seq"}.items():
        define(wb, name, f"{cq(ws.title)}${col}${LOG_FIRST}:${col}${LOG_LAST}")
    dv = DataValidation(type="date", operator="between", formula1="YearStart", formula2="DATE(2040,12,31)", allow_blank=True, error="أدخلي تاريخًا صحيحًا", errorTitle="تاريخ غير صحيح"); ws.add_data_validation(dv); dv.add(f"A{LOG_FIRST}:A{LOG_LAST}")
    dv2 = DataValidation(type="custom", formula1=f"COUNTIF(Stu_ID,B{LOG_FIRST})>0", allow_blank=True, error="الرقم الطلابي غير موجود في أوراق الصفوف", errorTitle="رقم غير معروف", prompt="اكتبي الرقم الطلابي كما هو في ورقة الصف", promptTitle="الرقم الطلابي"); ws.add_data_validation(dv2); dv2.add(f"B{LOG_FIRST}:B{LOG_LAST}")
    dv3 = DataValidation(type="list", formula1="=ContactMethods", allow_blank=True, error="اختاري طريقة من القائمة", errorTitle="قيمة غير صحيحة"); ws.add_data_validation(dv3); dv3.add(f"F{LOG_FIRST}:F{LOG_LAST}")
    tab = Table(displayName="tblContactLog", ref=f"A5:J{LOG_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.conditional_formatting.add(f"C{LOG_FIRST}:C{LOG_LAST}", FormulaRule(formula=[f'$C{LOG_FIRST}="رقم غير موجود"'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.conditional_formatting.add(f"A{LOG_FIRST}:H{LOG_LAST}", FormulaRule(formula=[f'AND($B{LOG_FIRST}<>"",OR($A{LOG_FIRST}="",$F{LOG_FIRST}=""))'], fill=fill(YEL_F)))
    ws.freeze_panes = "A6"
    protect(ws)
    return ws


def build_absence_log(wb):
    ws = wb.create_sheet(SHEETS["alog"])
    setup_sheet(ws, tab="9CA3AF")
    title_block(ws, "سجل الغياب", "يُجمَّع تلقائيًا من قاعدة البيانات: كل حالة غياب (بدون عذر أو بعذر) في سطر مستقل — استخدمي أسهم الفلترة في رأس الجدول", "Absence log (auto)")
    put(ws, "A3", "إجمالي حالات الغياب المسجَّلة:", f=font(10, True, NAVY), al=ALIGN_R)
    put(ws, "B3", f"=INDEX(DB_Seq,{DB_LAST-DB_FIRST+1})", f=font(11, True, TEAL), al=ALIGN_C)
    put(ws, "C3", f'=IF(B3>{ALOG_LAST-ALOG_FIRST+1},"⚠ تجاوز عدد الحالات سعة هذا السجل ("&{ALOG_LAST-ALOG_FIRST+1}&") – تُعرض أول "&{ALOG_LAST-ALOG_FIRST+1}&" حالة","")', f=font(9, True, RED_T), al=ALIGN_R)
    hdr = ["التاريخ", "اسم الطفل", "الرقم الطلابي", "الصف", "نوع الغياب", "إجمالي الغياب حتى هذا التاريخ", "مستوى التنبيه", "حالة التواصل (الحالية)", "الملاحظات", "مرجع"]
    header_row(ws, 5, 1, hdr, widths=[12, 22, 12, 9, 10, 12, 16, 18, 24, 3], height=36)
    for r in range(ALOG_FIRST, ALOG_LAST + 1):
        k = r - ALOG_FIRST + 1
        put(ws, f"J{r}", f'=IF({k}>$B$3,"",IFERROR(MATCH({k}-1,DB_Seq,1)+1,1))', f=font(8, False, MUTED), al=ALIGN_C)
        G = lambda expr: f'=IF($J{r}="","",{expr})'
        put(ws, f"A{r}", G(f"INDEX(DB_Date,$J{r})"), f=font(10), al=ALIGN_C, nf="dd/mm/yyyy")
        put(ws, f"B{r}", G(f"INDEX(DB_Name,$J{r})"), f=font(10, True), al=ALIGN_R)
        put(ws, f"C{r}", G(f"INDEX(DB_ID,$J{r})"), f=font(10), al=ALIGN_C)
        put(ws, f"D{r}", G(f"INDEX(DB_Class,$J{r})"), f=font(10), al=ALIGN_C)
        put(ws, f"E{r}", G(f"INDEX(DB_Status,$J{r})"), f=font(10), al=ALIGN_C)
        put(ws, f"F{r}", G(f"INDEX(DB_Cum,$J{r})"), f=font(10), al=ALIGN_C)
        put(ws, f"G{r}", G(f"INDEX(LevelLabels,MATCH(F{r},LevelValues,1))"), f=font(10), al=ALIGN_C)
        put(ws, f"H{r}", G(f'IFERROR(INDEX(Stu_ContactReq,MATCH(C{r},Stu_ID,0)),"")'), f=font(10), al=ALIGN_C)
        put(ws, f"I{r}", G(f'IFERROR(INDEX(Stu_Notes,MATCH(C{r},Stu_ID,0)),"")'), f=font(10), al=ALIGN_R)
    tab = Table(displayName="tblAbsenceLog", ref=f"A5:J{ALOG_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    rg = f"G{ALOG_FIRST}:G{ALOG_LAST}"
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($F{ALOG_FIRST}<>"",$F{ALOG_FIRST}>=Thr_Critical)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($F{ALOG_FIRST}<>"",$F{ALOG_FIRST}>=Thr_High)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($F{ALOG_FIRST}<>"",$F{ALOG_FIRST}>=Thr_Repeated)'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($F{ALOG_FIRST}<>"",$F{ALOG_FIRST}>=Thr_FollowUp)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add(f"E{ALOG_FIRST}:E{ALOG_LAST}", FormulaRule(formula=[f'$E{ALOG_FIRST}="غائب"'], font=Font(name=FONT, bold=True, color=ABS_T)))
    ws.conditional_formatting.add(f"E{ALOG_FIRST}:E{ALOG_LAST}", FormulaRule(formula=[f'$E{ALOG_FIRST}="بعذر"'], font=Font(name=FONT, color=EXC_T)))
    ws.conditional_formatting.add(f"H{ALOG_FIRST}:H{ALOG_LAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("مطلوب",$H{ALOG_FIRST}))'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.freeze_panes = "A6"
    protect(ws)
    return ws


def build_followup(wb):
    ws = wb.create_sheet(SHEETS["fu"])
    setup_sheet(ws, tab="D97706")
    title_block(ws, "قائمة المتابعة", "الأطفال الذين بلغ غيابهم حد المتابعة أو أكثر — مرتبون تلقائيًا حسب الأولوية (الأعلى غيابًا أولًا). جميع الصفوف.", "Follow-Up List (auto)")
    put(ws, "A3", "عدد الحالات:", f=font(10, True, NAVY), al=ALIGN_R)
    put(ws, "B3", "=COUNT(Stu_KeyFU)", f=font(11, True, TEAL), al=ALIGN_C)
    put(ws, "C3", "تحتاج تواصلًا:", f=font(10, True, NAVY), al=ALIGN_R)
    put(ws, "D3", '=COUNTIFS(Stu_KeyFU,">0",Stu_Contacted,"لا")', f=font(11, True, ORG_T), al=ALIGN_C)
    hdr = ["الأولوية", "اسم الطفل", "الصف", "إجمالي الغياب", "نسبة الحضور", "المستوى", "رقم ولي الأمر", "اسم ولي الأمر", "تم التواصل؟", "تاريخ آخر تواصل", "الإجراء المطلوب", "مرجع"]
    header_row(ws, 5, 1, hdr, widths=[8, 22, 9, 9, 9, 15, 13, 16, 10, 12, 40, 3], height=36)
    FU_LAST = 5 + 60
    for r in range(6, FU_LAST + 1):
        k = r - 5
        put(ws, f"L{r}", f'=IFERROR(MATCH(LARGE(Stu_KeyFU,{k}),Stu_KeyFU,0),"")', f=font(8, False, MUTED), al=ALIGN_C)
        G = lambda expr: f'=IF($L{r}="","",{expr})'
        put(ws, f"A{r}", G(str(k)), f=font(10, True, NAVY), al=ALIGN_C, b=box())
        put(ws, f"B{r}", G(f"INDEX(Stu_Name,$L{r})"), f=font(10, True), al=ALIGN_R, b=box())
        put(ws, f"C{r}", G(f"INDEX(Stu_Class,$L{r})"), f=font(10), al=ALIGN_C, b=box())
        put(ws, f"D{r}", G(f"INDEX(Stu_Abs,$L{r})"), f=font(10, True), al=ALIGN_C, b=box())
        put(ws, f"E{r}", G(f"INDEX(Stu_Rate,$L{r})"), f=font(10), al=ALIGN_C, b=box(), nf="0.0%")
        put(ws, f"F{r}", G(f"INDEX(Stu_Status,$L{r})"), f=font(10, True), al=ALIGN_C, b=box())
        put(ws, f"G{r}", G(f"INDEX(Stu_Phone,$L{r})"), f=font(10), al=ALIGN_C, b=box())
        put(ws, f"H{r}", G(f"INDEX(Stu_Guardian,$L{r})"), f=font(10), al=ALIGN_R, b=box())
        put(ws, f"I{r}", G(f"INDEX(Stu_Contacted,$L{r})"), f=font(10, True), al=ALIGN_C, b=box())
        put(ws, f"J{r}", G(f"INDEX(Stu_LastContact,$L{r})"), f=font(10), al=ALIGN_C, b=box(), nf="dd/mm/yyyy")
        put(ws, f"K{r}", G(f"INDEX(Stu_Action,$L{r})"), f=font(10), al=ALIGN_R, b=box())
        put(ws, f"M{r}", G(f"INDEX(Stu_Level,$L{r})"), f=font(8, False, MUTED), al=ALIGN_C)
    put(ws, "M5", "مستوى", f=font(8, True, WHITE), bg=MUTED, al=ALIGN_C)
    tab = Table(displayName="tblFollowUp", ref=f"A5:M{FU_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)
    rg = f"A6:K{FU_LAST}"
    ws.conditional_formatting.add(rg, FormulaRule(formula=['AND($M6<>"",$M6>=Thr_Critical)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T), stopIfTrue=True))
    ws.conditional_formatting.add(f"F6:F{FU_LAST}", FormulaRule(formula=['AND($M6<>"",$M6=Thr_High)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add(f"F6:F{FU_LAST}", FormulaRule(formula=['AND($M6<>"",$M6=Thr_Repeated)'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add(f"F6:F{FU_LAST}", FormulaRule(formula=['AND($M6<>"",$M6=Thr_FollowUp)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add(f"I6:I{FU_LAST}", FormulaRule(formula=['$I6="لا"'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add(f"I6:I{FU_LAST}", FormulaRule(formula=['$I6="نعم"'], fill=fill(GREEN_F), font=Font(name=FONT, bold=True, color=GREEN_T)))
    ws.freeze_panes = "A6"
    protect(ws)
    return ws
