# -*- coding: utf-8 -*-
from common import *
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import FormulaRule, CellIsRule

REG_HDR = ["م", "الرقم الطلابي", "اسم الطفل", "الغياب", "غياب بعذر", "تأخير", "نسبة الحضور", "حالة المتابعة", "التواصل مع ولي الأمر",
           "رقم التواصل", "اسم ولي الأمر", "ملاحظات",
           "أيام الدوام", "الحضور", "نسبة الغياب", "آخر تاريخ غياب", "تم التواصل؟", "تاريخ آخر تواصل", "طريقة آخر تواصل", "مستوى التنبيه", "بدء الاحتساب",
           "تاريخ بلوغ حد المتابعة", "تاريخ بلوغ الغياب المتكرر", "تاريخ بلوغ الغياب المرتفع", "تاريخ بلوغ الحالة الحرجة", "هاتف 1", "هاتف 2", "هاتف 3", "هاتف 4"]
REG_W = [4, 12, 30, 7, 7, 7, 9, 15, 17, 14, 22, 16, 6, 6, 7, 11, 8, 11, 12, 7, 11, 11, 11, 11, 11, 12, 12, 12, 12]
INPUT_COLS = {"B", "C", "J", "K", "L", "Z", "AA", "AB", "AC"}
HIDDEN_COLS = ["M", "N", "O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC"]

def build_register(wb, cname, idx):
    ws = wb.create_sheet(cname)
    setup_sheet(ws, grid=False, zoom=100, tab=TEAL if idx < ACTIVE_CLASSES else "9CA3AF")
    S = cname; C = RC
    RNG = lambda r: f"${DC0}{r}:${DC1}{r}"
    DATES_R = f"${DC0}$5:${DC1}$5"; SCH = f"${DC0}$6:${DC1}$6"; CONF = f"${DC0}$7:${DC1}$7"
    put(ws, "A1", "سجل الحضور اليومي للصف", f=font(16, True, NAVY), al=ALIGN_R)
    put(ws, "A2", "الصف:", f=font(10, True), al=ALIGN_R)
    input_cell(ws, "B2", cname)
    put(ws, "D2", "الأطفال:", f=font(10, True), al=ALIGN_R)
    put(ws, "E2", f'=COUNTIF($B${REG_FIRST}:$B${REG_LAST},"<>")', f=font(11, True, TEAL), al=ALIGN_C)
    put(ws, "F2", "الأيام المحصورة:", f=font(10, True), al=ALIGN_R)
    put(ws, "G2", f'=COUNTIF({CONF},"✓")', f=font(11, True, TEAL), al=ALIGN_C)
    put(ws, "H2", "نسبة حضور الصف:", f=font(10, True), al=ALIGN_R)
    put(ws, "I2", f'=IF(SUM(${DC0}${SUM_ROW["reg"]}:${DC1}${SUM_ROW["reg"]})=0,"",(SUM(${DC0}${SUM_ROW["pres"]}:${DC1}${SUM_ROW["pres"]})+IF(LateAsPresent="نعم",SUM(${DC0}${SUM_ROW["late"]}:${DC1}${SUM_ROW["late"]}),0))/SUM(${DC0}${SUM_ROW["reg"]}:${DC1}${SUM_ROW["reg"]}))', f=font(11, True, TEAL), al=ALIGN_C, nf="0.0%")
    put(ws, "J2", "الحصر مكتمل حتى:", f=font(10, True, NAVY), al=ALIGN_R)
    input_cell(ws, "K2", None, nf="dd/mm/yyyy")
    put(ws, "L2", '=IF($K$2="",MIN(TODAY(),YearEnd),$K$2)', f=font(10, True, TEAL), al=ALIGN_C, nf='dd/mm/yyyy')
    put(ws, "L1", "يُحتسب حتى (تلقائي):", f=font(8, False, MUTED), al=ALIGN_C)
    dvc = DataValidation(type="list", formula1="=Cal_Date", allow_blank=True, error="اختاري يوم دوام من القائمة أو اتركي الخلية فارغة", errorTitle="تاريخ غير صحيح",
                         prompt="كل يوم دوام حتى هذا التاريخ يُعتبر محصورًا: من لم يُسجَّل له غياب يُعتبر حاضرًا. فارغ = حتى اليوم", promptTitle="الحصر مكتمل حتى")
    ws.add_data_validation(dvc); dvc.add(ws["K2"])
    put(ws, "A3", "الحصر: كل الأطفال حاضرون افتراضيًا. في عمود اليوم اختاري من القائمة المنسدلة «غياب بعذر» أو «غياب بدون عذر» لمن غاب فقط (و«متأخر» عند الحاجة). الأعمدة الرمادية عطلات.", f=font(9, False, MUTED, True), al=ALIGN_R)
    put(ws, "H3", "الانتقال إلى تاريخ:", f=font(10, True, NAVY), al=ALIGN_R)
    input_cell(ws, "I3", None, nf="dd/mm/yyyy")
    put(ws, "J3", f'=IF($I$3="","⬅ اختاري تاريخًا",IFERROR(HYPERLINK("#{q(S)}!"&ADDRESS({REG_FIRST},{DATE_COL0-1}+MATCH($I$3,{DATES_R},0)),"⬅ فتح عمود "&TEXT($I$3,"dd/mm/yyyy")),"⚠ ليس يوم دوام"))', f=font(10, True, "1D4ED8"), al=ALIGN_R)
    dvd = DataValidation(type="list", formula1="=Cal_Date", allow_blank=True, error="اختاري يوم دوام من القائمة", errorTitle="تاريخ غير صحيح", prompt="للإدخال بأثر رجعي: اختاري التاريخ ثم اضغطي الرابط المجاور", promptTitle="تاريخ سابق")
    ws.add_data_validation(dvd); dvd.add(ws["I3"])
    put(ws, "K3", f'=HYPERLINK("#{q(S)}!"&ADDRESS({REG_FIRST},{DATE_COL0-1}+IFERROR(MATCH(TODAY(),{DATES_R},1),1)),"⬅ عمود اليوم")', f=font(10, True, "1D4ED8"), al=ALIGN_R)
    put(ws, "L3", f'=HYPERLINK("#{q(S)}!"&ADDRESS({REG_FIRST},{DATE_COL0-1}+IFERROR(MATCH($L$2,{DATES_R},1),1)),"⬅ آخر يوم محصور")', f=font(10, True, "1D4ED8"), al=ALIGN_R)
    lab_col = L(DATE_COL0 - 1)
    for r, t in [(4, "الشهر"), (5, "التاريخ"), (6, "يوم دراسي"), (7, "محصور ✓")]:
        put(ws, f"{lab_col}{r}", t, f=font(9, True, NAVY), al=ALIGN_R)
    header_row(ws, 8, 1, REG_HDR, widths=REG_W, height=42)
    for col in HIDDEN_COLS:
        ws.column_dimensions[col].hidden = True
        ws[f"{col}8"].fill = fill(MUTED)
    for i, d in enumerate(DATES):
        c = DATE_COL0 + i; col = L(c)
        ws.column_dimensions[col].width = 12.5
        put(ws, (4, c), AR_MONTHS[d.month - 1] if (i == 0 or DATES[i-1].month != d.month) else "", f=font(8, True, NAVY), al=Alignment(horizontal="right"))
        put(ws, (5, c), d, nf="dd/mm", f=font(9, True), al=ALIGN_C, b=box())
        put(ws, (6, c), f"={cq(SHEETS['cal'])}$H${CAL_FIRST + i}", nf='"";"";"عطلة"', f=font(8, False, MUTED), al=ALIGN_C)
        put(ws, (7, c), f'=IF(AND({col}$6=1,{col}$5<={CUTOFF_CELL}),"✓","")', f=font(9, True, GREEN_T), al=ALIGN_C)
        put(ws, (8, c), AR_DAYS[d.weekday()], f=font(8, True, WHITE), bg=NAVY, al=ALIGN_C, b=box(NAVY))
    ws.row_dimensions[4].height = 14; ws.row_dimensions[5].height = 18; ws.row_dimensions[6].height = 14; ws.row_dimensions[7].height = 16
    for r in range(REG_FIRST, REG_LAST + 1):
        G0 = lambda expr: f'=IF($B{r}="","",{expr})'
        crit = f'{SCH},1,{DATES_R},">="&$U{r}'
        f = {}
        f["A"] = f'=IF($B{r}="","",ROW()-{REG_FIRST-1})'
        f["U"] = f'=IF($B{r}="","",YearStart)'
        f["D"] = G0(f'COUNTIFS({RNG(r)},"غياب بدون عذر",{crit})')
        f["E"] = G0(f'COUNTIFS({RNG(r)},"غياب بعذر",{crit})')
        f["F"] = G0(f'COUNTIFS({RNG(r)},"متأخر",{crit})')
        f["N"] = G0(f'COUNTIFS({RNG(r)},"حاضر",{crit})+COUNTIFS({CONF},"✓",{crit})-COUNTIFS({RNG(r)},"<>",{CONF},"✓",{crit})')
        f["M"] = G0(f'$N{r}+$D{r}+$E{r}+$F{r}')
        f["G"] = f'=IF($M{r}="","",IF($M{r}=0,"",($N{r}+IF(LateAsPresent="نعم",$F{r},0))/$M{r}))'
        f["O"] = f'=IF($M{r}="","",IF($M{r}=0,"",($D{r}+$E{r})/$M{r}))'
        f["T"] = f'=IF($D{r}="","",IF($D{r}>=Thr_Critical,Thr_Critical,IF($D{r}>=Thr_High,Thr_High,IF($D{r}>=Thr_Repeated,Thr_Repeated,IF($D{r}>=Thr_FollowUp,Thr_FollowUp,0)))))'
        f["H"] = f'=IF($B{r}="","",INDEX(LevelLabels,MATCH($T{r},LevelValues,0)))'
        f["P"] = f'=IF($D{r}="","",IF($D{r}=0,"",_xlfn.MAXIFS({DATES_R},{RNG(r)},"غياب بدون عذر",{crit})))'
        f["R"] = f'=IF($B{r}="","",IF(_xlfn.MAXIFS(Log_Date,Log_ID,$B{r})=0,"",_xlfn.MAXIFS(Log_Date,Log_ID,$B{r})))'
        f["S"] = f'=IF($R{r}="","",IFERROR(INDEX(Log_Method,MATCH($B{r}&"|"&$R{r},Log_Key,0)),""))'
        f["Q"] = f'=IF($T{r}="","",IF($T{r}=0,"—",IF(AND($R{r}<>"",$R{r}>=IF($T{r}=Thr_Critical,$Y{r},IF($T{r}=Thr_High,$X{r},IF($T{r}=Thr_Repeated,$W{r},$V{r})))),"نعم","لا")))'
        f["I"] = f'=IF($T{r}="","",IF($T{r}=0,"—",IF($Q{r}="نعم","✓ تم التواصل","⚠ التواصل مطلوب")))'
        for col, fm in f.items():
            put(ws, f"{col}{r}", fm, f=font(10), al=ALIGN_C, b=box())
        for col, thr in zip(["V", "W", "X", "Y"], ["Thr_FollowUp", "Thr_Repeated", "Thr_High", "Thr_Critical"]):
            ws[f"{col}{r}"] = ArrayFormula(f"{col}{r}", f'=IF($D{r}="","",IFERROR(SMALL(IF(({RNG(r)}="غياب بدون عذر")*({SCH}=1)*({DATES_R}>=$U{r}),{DATES_R}),{thr}),""))')
            style(ws[f"{col}{r}"], f=font(10), al=ALIGN_C, b=box(), nf="dd/mm/yyyy")
        for col in ["P", "R", "U"]:
            ws[f"{col}{r}"].number_format = "dd/mm/yyyy"
        ws[f"G{r}"].number_format = "0.0%"; ws[f"O{r}"].number_format = "0.0%"
        for col in ["C", "K", "L"]:
            ws[f"{col}{r}"].alignment = ALIGN_R
        for col in INPUT_COLS:
            c = ws[f"{col}{r}"]
            c.protection = UNLOCKED; c.fill = fill(INPUT_FILL); c.border = box(INPUT_BORDER); c.font = font(10, False, TXT)
        ws[f"C{r}"].font = font(10, True, TXT)
        # phone dropdown: the visible cell lists that child's numbers from the hidden block Z:AC
        dvp = DataValidation(type="list", formula1=f"=$Z${r}:$AC${r}", allow_blank=True, showErrorMessage=False, prompt="اختاري الرقم من القائمة (أرقام ولي الأمر المسجَّلة)", promptTitle="رقم التواصل")
        dvp._keep_soft = True
        ws.add_data_validation(dvp); dvp.add(ws[f"J{r}"])
        for i in range(NDAYS):
            c = ws.cell(r, DATE_COL0 + i)
            c.protection = UNLOCKED; c.font = font(8); c.alignment = ALIGN_C; c.border = box("EEF0F3")
        ws.row_dimensions[r].height = 20
    labels = {"abs": "غياب بدون عذر", "exc": "غياب بعذر", "late": "متأخر", "pres": "حاضر", "reg": "المحصورون", "rate": "نسبة الحضور"}
    put(ws, f"A{SUM_ROW['abs']-1}", "ملخص اليوم (تلقائي)", f=font(10, True, NAVY), al=ALIGN_R)
    for key, r in SUM_ROW.items():
        put(ws, f"C{r}", labels[key], f=font(10, True, NAVY), al=ALIGN_R, bg=LIGHT)
        put(ws, f"{lab_col}{r}", labels[key], f=font(9, True, NAVY), al=ALIGN_R)
        for i in range(NDAYS):
            c = DATE_COL0 + i; col = L(c)
            base = f'$U${REG_FIRST}:$U${REG_LAST},"<="&{col}$5'
            stu = f'{col}${REG_FIRST}:{col}${REG_LAST}'
            if key == "abs": fm = f'=IF({col}$6=1,COUNTIFS({stu},"غياب بدون عذر",{base}),0)'
            elif key == "exc": fm = f'=IF({col}$6=1,COUNTIFS({stu},"غياب بعذر",{base}),0)'
            elif key == "late": fm = f'=IF({col}$6=1,COUNTIFS({stu},"متأخر",{base}),0)'
            elif key == "pres": fm = f'=IF({col}$6=1,COUNTIFS({stu},"حاضر",{base})+IF({col}$7="✓",COUNTIFS($B${REG_FIRST}:$B${REG_LAST},"<>",{base})-COUNTIFS({stu},"<>",{base}),0),0)'
            elif key == "reg": fm = f'={col}{SUM_ROW["abs"]}+{col}{SUM_ROW["exc"]}+{col}{SUM_ROW["late"]}+{col}{SUM_ROW["pres"]}'
            else: fm = f'=IF({col}{SUM_ROW["reg"]}=0,"",({col}{SUM_ROW["pres"]}+IF(LateAsPresent="نعم",{col}{SUM_ROW["late"]},0))/{col}{SUM_ROW["reg"]})'
            put(ws, (r, c), fm, f=font(9, True if key in ("reg", "rate") else False, NAVY if key in ("reg", "rate") else TXT), al=ALIGN_C, bg=LIGHT, nf="0%" if key == "rate" else "0")
    dv = DataValidation(type="list", formula1='"غياب بعذر,غياب بدون عذر,حاضر,متأخر"', allow_blank=True, error="اختاري من القائمة: غياب بعذر / غياب بدون عذر / حاضر / متأخر", errorTitle="حالة غير صحيحة", prompt="الجميع حاضر افتراضيًا – اختاري «غياب بعذر» أو «غياب بدون عذر» لمن غاب", promptTitle="حالة الحضور")
    ws.add_data_validation(dv); dv.add(f"{DC0}{REG_FIRST}:{DC1}{REG_LAST}")
    dv4 = DataValidation(type="custom", formula1=f"COUNTIF(Stu_ID,B{REG_FIRST})<=1", allow_blank=True, error="هذا الرقم الطلابي مسجَّل مسبقًا لطفل آخر", errorTitle="رقم طلابي مكرر", prompt="رقم فريد لكل طفل (يُكتب مرة واحدة)", promptTitle="الرقم الطلابي")
    ws.add_data_validation(dv4); dv4.add(f"B{REG_FIRST}:B{REG_LAST}")
    dv6 = DataValidation(type="list", formula1="=ClassList", allow_blank=False, error="اختاري صفًا من قائمة الإعدادات", errorTitle="صف غير معروف")
    ws.add_data_validation(dv6); dv6.add(ws["B2"])
    grid = f"{DC0}{REG_FIRST}:{DC1}{REG_LAST}"
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f"{DC0}$6=0"], fill=fill(GRAY_F), stopIfTrue=True))
    ws.conditional_formatting.add(grid, CellIsRule(operator="equal", formula=['"غياب بدون عذر"'], fill=fill(ABS_F), font=Font(name=FONT, bold=True, color=ABS_T)))
    ws.conditional_formatting.add(grid, CellIsRule(operator="equal", formula=['"غياب بعذر"'], fill=fill(EXC_F), font=Font(name=FONT, color=EXC_T)))
    ws.conditional_formatting.add(grid, CellIsRule(operator="equal", formula=['"متأخر"'], fill=fill(LATE_F), font=Font(name=FONT, color=LATE_T)))
    ws.conditional_formatting.add(grid, CellIsRule(operator="equal", formula=['"حاضر"'], font=Font(name=FONT, color=GREEN_T)))
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'AND({DC0}$7="",{DC0}$6=1,{DC0}{REG_FIRST}="",$B{REG_FIRST}<>"")'], fill=fill("FAFAFA"), font=Font(name=FONT, color="C0C4CC")))
    ws.conditional_formatting.add(f"{DC0}4:{DC1}8", FormulaRule(formula=[f"{DC0}$5=TODAY()"], fill=fill(BLUE_F)))
    ws.conditional_formatting.add(f"{DC0}4:{DC1}8", FormulaRule(formula=[f"{DC0}$5={CUTOFF_CELL}"], fill=fill(TEAL_L)))
    ws.conditional_formatting.add(f"{DC0}7:{DC1}7", FormulaRule(formula=[f'AND({DC0}$7="",{DC0}$6=1,COUNTA({DC0}${REG_FIRST}:{DC0}${REG_LAST})>0)'], fill=fill(YEL_F), font=Font(name=FONT, color=YEL_T)))
    ws.conditional_formatting.add(f"{DC0}5:{DC1}6", FormulaRule(formula=[f"{DC0}$6=0"], fill=fill(GRAY_F), font=Font(name=FONT, color=MUTED)))
    lvl = f"A{REG_FIRST}:I{REG_LAST}"
    ws.conditional_formatting.add(lvl, FormulaRule(formula=[f'AND($T{REG_FIRST}<>"",$T{REG_FIRST}>=Thr_Critical)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T), border=Border(left=side(RED_T, "medium"), right=side(RED_T, "medium")), stopIfTrue=True))
    for col in ["D", "H"]:
        rg = f"{col}{REG_FIRST}:{col}{REG_LAST}"
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($T{REG_FIRST}<>"",$T{REG_FIRST}=Thr_High)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($T{REG_FIRST}<>"",$T{REG_FIRST}=Thr_Repeated)'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($T{REG_FIRST}<>"",$T{REG_FIRST}=Thr_FollowUp)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND($T{REG_FIRST}<>"",$T{REG_FIRST}=0)'], fill=fill(GREEN_F), font=Font(name=FONT, color=GREEN_T)))
    ws.conditional_formatting.add(f"I{REG_FIRST}:I{REG_LAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("مطلوب",$I{REG_FIRST}))'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add(f"B{REG_FIRST}:B{REG_LAST}", FormulaRule(formula=[f'AND($B{REG_FIRST}<>"",COUNTIF(Stu_ID,$B{REG_FIRST})>1)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    ws.freeze_panes = f"D{REG_FIRST}"
    if idx >= ACTIVE_CLASSES:
        ws.sheet_state = "hidden"
    protect(ws)
    return ws
