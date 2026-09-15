# -*- coding: utf-8 -*-
from common import *
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.page import PageMargins
from dashboard import list_header

def build_report(wb):
    ws = wb.create_sheet(SHEETS["rpt"])
    setup_sheet(ws, tab=NAVY)
    for c, w in zip("ABCDEFGH", [22, 11, 11, 11, 11, 12, 14, 18]):
        ws.column_dimensions[c].width = w
    put(ws, "A1", "=SchoolName", f=font(16, True, NAVY), al=ALIGN_R)
    put(ws, "A2", '="قسم "&SectionName&"   |   العام الدراسي "&AcademicYear', f=font(10, False, TEAL), al=ALIGN_R)
    put(ws, "E1", "KG1 Attendance Summary", f=font(14, True, NAVY), al=Alignment(horizontal="left"))
    put(ws, "E2", "ملخص الحضور والغياب", f=font(11, False, TEAL), al=Alignment(horizontal="left"))
    put(ws, "A4", "نوع الفترة", f=font(9, True, NAVY), al=ALIGN_R); input_cell(ws, "B4", "الشهر")
    put(ws, "C4", "التاريخ", f=font(9, True, NAVY), al=ALIGN_R); input_cell(ws, "D4", None, nf="dd/mm/yyyy")
    put(ws, "E4", "الصف", f=font(9, True, NAVY), al=ALIGN_R); input_cell(ws, "F4", "الكل")
    put(ws, "G4", '="تاريخ الطباعة: "&TEXT(TODAY(),"dd/mm/yyyy")', f=font(8, False, MUTED), al=ALIGN_R)
    define(wb, "Rpt_Period", f"{cq(ws.title)}$B$4"); define(wb, "Rpt_DateInput", f"{cq(ws.title)}$D$4"); define(wb, "Rpt_Class", f"{cq(ws.title)}$F$4")
    dv = DataValidation(type="list", formula1="=RptPeriodList", allow_blank=False); ws.add_data_validation(dv); dv.add(ws["B4"])
    dv2 = DataValidation(type="list", formula1="=FilterClassList", allow_blank=False); ws.add_data_validation(dv2); dv2.add(ws["F4"])
    dv3 = DataValidation(type="date", operator="between", formula1="YearStart", formula2="YearEnd", allow_blank=True, error="أدخلي تاريخًا ضمن العام الدراسي", errorTitle="تاريخ غير صحيح"); ws.add_data_validation(dv3); dv3.add(ws["D4"])
    # hidden calc (column J/K – outside print area)
    ws.column_dimensions["J"].width = 20; ws.column_dimensions["K"].width = 14
    calc = [("Rpt_Date", "التاريخ الفعّال", '=IF(Rpt_DateInput="",LastRecorded,Rpt_DateInput)', "dd/mm/yyyy"),
            ("Rpt_PStart", "بداية الفترة", '=IF(Rpt_Period="يوم",Rpt_Date,IF(Rpt_Period="الشهر",DATE(YEAR(Rpt_Date),MONTH(Rpt_Date),1),IF(Rpt_Period="الفصل الدراسي",IFERROR(INDEX(TermStart,MATCH(Rpt_Date,TermStart,1)),YearStart),YearStart)))', "dd/mm/yyyy"),
            ("Rpt_PEnd", "نهاية الفترة", '=IF(Rpt_Period="يوم",Rpt_Date,IF(Rpt_Period="الشهر",EOMONTH(Rpt_Date,0),IF(Rpt_Period="الفصل الدراسي",IFERROR(INDEX(TermEnd,MATCH(Rpt_Date,TermStart,1)),YearEnd),YearEnd)))', "dd/mm/yyyy")]
    for i, (name, lab, fm, nf) in enumerate(calc):
        put(ws, f"J{2+i}", lab, f=font(8, False, MUTED), al=ALIGN_R); put(ws, f"K{2+i}", fm, f=font(8, False, MUTED), al=ALIGN_C, nf=nf)
        define(wb, name, f"{cq(ws.title)}$K${2+i}")
    put(ws, "A5", '="الفترة: "&Rpt_Period&"  ("&TEXT(Rpt_PStart,"dd/mm/yyyy")&" – "&TEXT(Rpt_PEnd,"dd/mm/yyyy")&")   |   الصف: "&Rpt_Class', f=font(10, True, TEAL), al=ALIGN_R)
    ws.row_dimensions[5].height = 22
    # summary cards
    cards = [("عدد الأطفال", "=SUM(Stu_RptIn)", None), ("الأيام المسجَّلة في الفترة", '=COUNTIFS(Cal_Date,">="&Rpt_PStart,Cal_Date,"<="&Rpt_PEnd,Cal_Reg,">0")', None),
             ("نسبة الحضور", '=IF(SUM(Stu_RTot)=0,"—",SUM(Stu_RAtt)/SUM(Stu_RTot))', "0.0%"), ("إجمالي حالات الغياب في الفترة", "=SUM(Stu_RAbs)", None),
             ('="حالات المتابعة ("&Thr_FollowUp&"+)"', '=COUNTIFS(Stu_RptIn,1,Stu_Level,">="&Thr_FollowUp)', None), ('="حالات حرجة ("&Thr_Critical&"+)"', '=COUNTIFS(Stu_RptIn,1,Stu_Level,">="&Thr_Critical)', None)]
    for j, (lab, fm, nf) in enumerate(cards):
        r = 7 + (j // 3) * 3; c = 1 + (j % 3) * 3
        for rr in (r, r + 1):
            for cc in range(c, c + 2):
                ws.cell(rr, cc).fill = fill(LIGHT); ws.cell(rr, cc).border = box(WHITE)
        put(ws, (r, c), lab, f=font(9, True, MUTED), al=Alignment(horizontal="centerContinuous", vertical="center")); ws.cell(r, c + 1).alignment = Alignment(horizontal="centerContinuous")
        put(ws, (r + 1, c), fm, f=font(16, True, NAVY), al=Alignment(horizontal="centerContinuous", vertical="center"), nf=nf); ws.cell(r + 1, c + 1).alignment = Alignment(horizontal="centerContinuous")
        ws.row_dimensions[r + 1].height = 26
    # class table
    section(ws, "A13", "ملخص الصفوف في الفترة", None, 8)
    list_header(ws, 14, 1, ["الصف", "عدد الأطفال", "نسبة الحضور", "الغياب (بدون عذر)", "غياب بعذر", '="حالات "&Thr_FollowUp&"+"', '="حالات "&Thr_Critical&"+"'])
    for k in range(NCLS):
        r = 15 + k
        put(ws, f"A{r}", f'=IF(INDEX(ClassActive,{k+1})<>"نعم","",IF(OR(Rpt_Class="الكل",Rpt_Class=INDEX(ClassList,{k+1})),INDEX(ClassList,{k+1}),""))', f=font(10, True), al=ALIGN_C, b=bottom())
        G = lambda e: f'=IF($A{r}="","",{e})'
        put(ws, f"B{r}", G(f'COUNTIFS(Stu_Class,$A{r},Stu_StStatus,"نشط")'), f=font(10), al=ALIGN_C, b=bottom())
        P = f'COUNTIFS(DB_Class,$A{r},DB_Date,">="&Rpt_PStart,DB_Date,"<="&Rpt_PEnd,DB_Code,"P")'; Lt = P.replace('"P"', '"L"'); A = P.replace('"P"', '"A"'); E = P.replace('"P"', '"E"')
        put(ws, f"C{r}", G(f'IF(({P}+{A}+{E}+{Lt})=0,"—",({P}+IF(LateAsPresent="نعم",{Lt},0))/({P}+{A}+{E}+{Lt}))'), f=font(10, True, TEAL), al=ALIGN_C, b=bottom(), nf="0.0%")
        put(ws, f"D{r}", G(A), f=font(10), al=ALIGN_C, b=bottom()); put(ws, f"E{r}", G(E), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"F{r}", G(f'COUNTIFS(Stu_Class,$A{r},Stu_StStatus,"نشط",Stu_Level,">="&Thr_FollowUp)'), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"G{r}", G(f'COUNTIFS(Stu_Class,$A{r},Stu_StStatus,"نشط",Stu_Level,">="&Thr_Critical)'), f=font(10), al=ALIGN_C, b=bottom())
    ws.conditional_formatting.add("G15:G21", FormulaRule(formula=['AND($G15<>"",$G15>0)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T)))
    # cases
    section(ws, "A24", "أهم الحالات التي تحتاج إلى متابعة (مرتبة حسب إجمالي الغياب)", None, 8)
    list_header(ws, 25, 1, ["اسم الطفل", "الصف", "إجمالي الغياب", "نسبة الحضور", "المستوى", "تم التواصل؟", "الإجراء المطلوب"])
    ws.merge_cells("G25:H25")
    for k in range(1, 16):
        r = 25 + k
        idx = f"IFERROR(MATCH(LARGE(Stu_KeyRpt,{k}),Stu_KeyRpt,0),\"\")"
        put(ws, f"J{r}", f"={idx}", f=font(8, False, MUTED), al=ALIGN_C)
        G = lambda e: f'=IF($J{r}="","",{e})'
        put(ws, f"A{r}", G(f"INDEX(Stu_Name,$J{r})"), f=font(10, True), al=ALIGN_R, b=bottom())
        put(ws, f"B{r}", G(f"INDEX(Stu_Class,$J{r})"), f=font(10), al=ALIGN_C, b=bottom())
        put(ws, f"C{r}", G(f"INDEX(Stu_Abs,$J{r})"), f=font(10, True), al=ALIGN_C, b=bottom())
        put(ws, f"D{r}", G(f"INDEX(Stu_Rate,$J{r})"), f=font(10), al=ALIGN_C, b=bottom(), nf="0.0%")
        put(ws, f"E{r}", G(f"INDEX(Stu_Status,$J{r})"), f=font(10, True), al=ALIGN_C, b=bottom())
        put(ws, f"F{r}", G(f"INDEX(Stu_Contacted,$J{r})"), f=font(10, True), al=ALIGN_C, b=bottom())
        put(ws, f"G{r}", G(f"INDEX(Stu_Action,$J{r})"), f=font(9), al=ALIGN_R, b=bottom()); ws.merge_cells(f"G{r}:H{r}")
        put(ws, f"K{r}", G(f"INDEX(Stu_Level,$J{r})"), f=font(8, False, MUTED), al=ALIGN_C)
    ws.conditional_formatting.add("A26:H40", FormulaRule(formula=['AND($K26<>"",$K26>=Thr_Critical)'], fill=fill(RED_F), font=Font(name=FONT, bold=True, color=RED_T), stopIfTrue=True))
    ws.conditional_formatting.add("E26:E40", FormulaRule(formula=['AND($K26<>"",$K26=Thr_High)'], fill=fill(DORG_F), font=Font(name=FONT, bold=True, color=DORG_T)))
    ws.conditional_formatting.add("E26:E40", FormulaRule(formula=['AND($K26<>"",$K26=Thr_Repeated)'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    ws.conditional_formatting.add("E26:E40", FormulaRule(formula=['AND($K26<>"",$K26=Thr_FollowUp)'], fill=fill(YEL_F), font=Font(name=FONT, bold=True, color=YEL_T)))
    ws.conditional_formatting.add("F26:F40", FormulaRule(formula=['$F26="لا"'], fill=fill(ORG_F), font=Font(name=FONT, bold=True, color=ORG_T)))
    put(ws, "A41", '=IF(COUNT(Stu_KeyRpt)=0,"لا توجد حالات تحتاج إلى متابعة ✓",IF(COUNT(Stu_KeyRpt)>15,"تُعرض أول 15 حالة من أصل "&COUNT(Stu_KeyRpt)&" – القائمة الكاملة في ورقة «قائمة المتابعة»",""))', f=font(9, False, MUTED, True), al=ALIGN_R)
    put(ws, "A44", "المسجّلة: ______________________", f=font(10), al=ALIGN_R)
    put(ws, "E44", "مديرة الروضة: ______________________", f=font(10), al=ALIGN_R)
    ws.print_area = "A1:H45"
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "&A – &D"
    protect(ws)
    return ws


GUIDE = [
    ("h", "دليل الاستخدام السريع"),
    ("p", "هذا الملف نظام متكامل لمتابعة وتحليل الحضور والغياب لقسم الروضة الأولى KG1. الإدخال اليومي يتم في ورقة كل صف، وكل ما عداه (لوحة المتابعة، القوائم، السجلات، التقارير) يُحدَّث تلقائيًا."),
    ("h2", "الإعداد مرة واحدة في بداية العام"),
    ("s", "1. افتحي ورقة «الإعدادات» وراجعي اسم المدرسة والعام الدراسي وحدود التنبيه (3 / 5 / 10 / 15) والفصول الدراسية والعطلات الرسمية."),
    ("s", "2. في ورقة كل صف (KG1-1 … KG1-7) اكتبي الأطفال مرة واحدة: الرقم الطلابي، اسم الطفل، الحالة الدراسية (نشط)، اسم ولي الأمر ورقم التواصل. لا تكرري الأسماء في أي ورقة أخرى."),
    ("s", "3. طفل التحق متأخرًا؟ اكتبي «تاريخ الالتحاق» فلا تُحتسب الأيام السابقة له. طفل انسحب أو انتقل؟ غيّري «الحالة الدراسية» إلى منسحب/منقول ولا تحذفي صفه."),
    ("s", "4. الصفوف السبعة KG1-1 … KG1-7 جاهزة. لإيقاف صف (دمج أو إغلاق): في «الإعدادات» غيّري «مفعّل» إلى لا وأخفي ورقته؛ يمكن إعادة تفعيله في أي وقت دون فقد البيانات."),
    ("h2", "الاستخدام اليومي (أقل من دقيقة لكل صف)"),
    ("s", "الخطوة 1: افتحي ورقة الصف واضغطي رابط «الانتقال إلى عمود اليوم» أعلى الورقة (أو استخدمي ورقة «الحضور اليومي» للانتقال إلى أي تاريخ)."),
    ("s", "الخطوة 2: في عمود اليوم اكتبي حالة الغائبين والمتأخرين فقط من القائمة المنسدلة: غائب / بعذر / متأخر (يمكن كتابة «حاضر» أيضًا لكن لا حاجة لذلك)."),
    ("s", "الخطوة 3: ضعي ✓ في صف «تأكيد تسجيل اليوم» أعلى العمود. عندها فقط يُعتبر كل طفل لم تُكتب له حالة حاضرًا. بدون ✓ لا يُحتسب اليوم لهؤلاء الأطفال (الخلية الفارغة لا تعني غيابًا ولا حضورًا)."),
    ("s", "الخطوة 4: احفظي الملف، ثم افتحي «لوحة المتابعة» لمشاهدة نسب الحضور وتنبيهات اليوم ومن وصل إلى حد التنبيه."),
    ("s", "الخطوة 5: تواصلي مع أولياء الأمور المطلوبين، وسجّلي كل تواصل في ورقة «سجل التواصل» (التاريخ + الرقم الطلابي + الطريقة + النتيجة). يتحول عمود «تم التواصل؟» إلى نعم تلقائيًا."),
    ("h2", "معنى الألوان ومستويات التنبيه"),
    ("s", "أخضر: منتظم (أقل من 3 غيابات). أصفر: يحتاج متابعة (3+) – ذكّري ولي الأمر. برتقالي: غياب متكرر (5+) – ضرورة التواصل مع ولي الأمر. برتقالي داكن: غياب مرتفع (10+) – تدخل إداري مطلوب. أحمر: حالة حرجة (15+) – يظهر الطفل في قائمة «حالات الغياب الحرجة»."),
    ("s", "تُحتسب في حدود التنبيه أيام الغياب بدون عذر فقط. الغياب بعذر يُعرض في عمود مستقل ويؤثر على نسبة الحضور. المتأخر يُعتبر حاضرًا (يمكن تغيير ذلك في الإعدادات)."),
    ("s", "نسبة الحضور = (أيام الحضور + التأخير) ÷ الأيام المسجَّلة فعليًا للطفل × 100. لا تُحتسب العطلات ولا الأيام غير المسجَّلة ولا الأيام السابقة لتاريخ الالتحاق."),
    ("s", "«تم التواصل؟» = نعم إذا سُجِّل تواصل بتاريخ يساوي أو يلي تاريخ بلوغ الطفل مستواه الحالي؛ فإذا ارتفع مستواه بعد التواصل يعود إلى «لا» ليُطلب تواصل جديد."),
    ("h2", "الأوراق"),
    ("s", "لوحة المتابعة: مؤشرات اليوم، التنبيهات اليومية، نسبة الحضور العامة حسب الفترة (اليوم/الأسبوع/الشهر/الفصل/العام)، مقارنة الصفوف، الأعلى حضورًا، من يحتاج دعمًا، الحالات الحرجة. اختيار صف معين يحدّث كل الأرقام."),
    ("s", "الحضور اليومي: مراجعة حالة أي يوم لأي صف والانتقال السريع إلى عمود التسجيل. | قائمة المتابعة: كل من بلغ 3 غيابات أو أكثر مرتبين حسب الأولوية. | سجل الغياب: كل حالة غياب في سطر مستقل مع الفلاتر. | ملف الطفل: اختيار اسم لعرض ملفه وتقويم حضوره الشهري."),
    ("s", "الاتجاهات: التحليل اليومي والأسبوعي والشهري والفصلي. | التقرير: ملخص A4 للطباعة (يوم / شهر / فصل / عام – لكل الصفوف أو صف محدد). | الطلاب + التقويم + قاعدة البيانات: أوراق تلقائية للتحليل، لا تحتاج إلى إدخال."),
    ("h2", "الحماية والأخطاء الشائعة"),
    ("s", "الأوراق محمية بدون كلمة مرور؛ الخلايا الصفراء فقط قابلة للتعديل. لإلغاء الحماية عند الحاجة: مراجعة ← إلغاء حماية الورقة."),
    ("s", "لا تُدرجي ولا تحذفي صفوفًا أو أعمدة داخل أوراق الصفوف؛ استخدمي الصفوف الثلاثين الجاهزة بالترتيب، وللطفل المنسحب غيّري حالته فقط."),
    ("s", "الرقم الطلابي يجب أن يكون فريدًا؛ يرفض الملف الرقم المكرر ويلوّنه بالأحمر. حالة الحضور تُختار من القائمة فقط."),
    ("s", "عند فتح الملف في Excel قد يظهر «تمكين التحرير» أو «تمكين المحتوى» – اضغطيه ليتم حساب المعادلات. الملف لا يحتوي على وحدات ماكرو."),
]

def build_guide(wb):
    ws = wb.create_sheet(SHEETS["guide"])
    setup_sheet(ws, tab="9CA3AF")
    ws.column_dimensions["A"].width = 130
    r = 1
    for kind, text in GUIDE:
        if kind == "h":
            put(ws, f"A{r}", text, f=font(18, True, NAVY), al=ALIGN_R); ws.row_dimensions[r].height = 30
        elif kind == "h2":
            r += 1; put(ws, f"A{r}", text, f=font(12, True, TEAL), al=ALIGN_R, b=Border(bottom=side(TEAL, "medium"))); ws.row_dimensions[r].height = 24
        elif kind == "p":
            put(ws, f"A{r}", text, f=font(10), al=ALIGN_RW); ws.row_dimensions[r].height = 34
        else:
            put(ws, f"A{r}", text, f=font(10), al=ALIGN_RW); ws.row_dimensions[r].height = 18 * (1 + len(text) // 120)
        r += 1
    put(ws, f"A{r+1}", '=HYPERLINK("#\'لوحة المتابعة\'!A1","⬅ العودة إلى لوحة المتابعة")', f=font(11, True, "1D4ED8"), al=ALIGN_R)
    protect(ws)
    return ws
