# -*- coding: utf-8 -*-
from common import *
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
import datetime as dt

TERMS = [("الفصل الأول", dt.date(2026, 8, 31), dt.date(2026, 12, 11)),
         ("الفصل الثاني", dt.date(2027, 1, 4), dt.date(2027, 3, 26)),
         ("الفصل الثالث", dt.date(2027, 4, 12), dt.date(2027, 6, 30))]
HOLIDAYS = [(dt.date(2026, 12, 2), dt.date(2026, 12, 3), "اليوم الوطني (تقديري – يرجى التحديث)"),
            (dt.date(2026, 12, 14), dt.date(2027, 1, 1), "إجازة الشتاء (تقديري)"),
            (dt.date(2027, 3, 9), dt.date(2027, 3, 11), "عيد الفطر (تقديري)"),
            (dt.date(2027, 3, 29), dt.date(2027, 4, 9), "إجازة الربيع (تقديري)"),
            (dt.date(2027, 5, 16), dt.date(2027, 5, 19), "عيد الأضحى (تقديري)")]

def build_settings(wb):
    ws = wb.create_sheet(SHEETS["set"])
    setup_sheet(ws, tab="9CA3AF")
    title_block(ws, "الإعدادات", "القيم في الخلايا الصفراء قابلة للتعديل — تعتمد عليها جميع المعادلات والتنبيهات في الملف", "Settings")
    for col, w in zip("ABCDEFG", [30, 18, 30, 44, 14, 3, 18]):
        ws.column_dimensions[col].width = w
    r = 5
    section(ws, f"A{r}", "بيانات عامة", "General", 4); r += 1
    general = [("اسم المدرسة", "روضة الخليف", "SchoolName", None),
               ("النطاق التعليمي", "AD 2.3", "SchoolZone", None),
               ("القسم", "الروضة الأولى KG1", "SectionName", None),
               ("العام الدراسي", "2026–2027", "AcademicYear", None),
               ("الفصل الدراسي الحالي", "الفصل الأول", "CurrentTerm", "terms"),
               ("بداية العام الدراسي (ثابت – أُنشئت أعمدة السجل عليه)", YEAR_START, "YearStart", "date"),
               ("نهاية العام الدراسي (ثابت)", YEAR_END, "YearEnd", "date"),
               ("هل الجمعة يوم دراسي؟", "نعم", "FridaySchool", "yesno"),
               ("احتساب المتأخر ضمن الحضور؟", "نعم", "LateAsPresent", "yesno")]
    names = {}
    for lab, val, name, kind in general:
        put(ws, f"A{r}", lab, f=font(10), al=ALIGN_R)
        c = input_cell(ws, f"B{r}", val, nf="dd/mm/yyyy" if kind == "date" else None)
        if kind == "date":
            c.protection = Protection(locked=True); c.fill = fill(GRAY_F2); c.font = font(10, False, MUTED)
        names[name] = f"{cq(ws.title)}$B${r}"
        if kind == "yesno":
            names.setdefault("_yesno_cells", []).append(f"B{r}")
        if kind == "terms":
            names["_term_cell"] = f"B{r}"
        r += 1
    r += 1
    section(ws, f"A{r}", "حدود التنبيه (عدد أيام الغياب بدون عذر)", "Alert thresholds", 4); r += 1
    header_row(ws, r, 1, ["المستوى", "الحد (عدد الغيابات)", "حالة المتابعة (تظهر في الملف)", "الإجراء المطلوب"], height=24); r += 1
    levels = [("منتظم", 0, "منتظم", "—"),
              ("يحتاج متابعة", 3, "⚠ يحتاج متابعة", "التواصل مع ولي الأمر (تذكير)"),
              ("غياب متكرر", 5, "⚠ غياب متكرر", "ضرورة التواصل مع ولي الأمر"),
              ("غياب مرتفع", 10, "⚠ غياب مرتفع", "تدخل إداري مطلوب"),
              ("حالة حرجة", 15, "🔴 حالة حرجة", "حالة حرجة – غياب 15 يومًا أو أكثر – تدخل إداري عاجل")]
    lv0 = r
    lvl_fills = [(GREEN_F, GREEN_T), (YEL_F, YEL_T), (ORG_F, ORG_T), (DORG_F, DORG_T), (RED_F, RED_T)]
    for (lab, v, st, act), (bgc, fgc) in zip(levels, lvl_fills):
        put(ws, f"A{r}", lab, f=font(10, True, fgc), bg=bgc, al=ALIGN_R, b=box())
        if v == 0:
            put(ws, f"B{r}", 0, f=font(10, False, MUTED), al=ALIGN_C, b=box())
        else:
            input_cell(ws, f"B{r}", v)
        put(ws, f"C{r}", st, f=font(10), al=ALIGN_R, b=box())
        put(ws, f"D{r}", act, f=font(10), al=ALIGN_RW, b=box())
        r += 1
    names["Thr_FollowUp"] = f"{cq(ws.title)}$B${lv0+1}"
    names["Thr_Repeated"] = f"{cq(ws.title)}$B${lv0+2}"
    names["Thr_High"] = f"{cq(ws.title)}$B${lv0+3}"
    names["Thr_Critical"] = f"{cq(ws.title)}$B${lv0+4}"
    names["LevelValues"] = f"{cq(ws.title)}$B${lv0}:$B${lv0+4}"
    names["LevelLabels"] = f"{cq(ws.title)}$C${lv0}:$C${lv0+4}"
    names["LevelActions"] = f"{cq(ws.title)}$D${lv0}:$D${lv0+4}"
    put(ws, f"A{r}", "ملاحظة: يُحتسب في حدود التنبيه الغياب بدون عذر فقط؛ الغياب بعذر يُعرض في عمود مستقل ويؤثر على نسبة الحضور.", f=font(9, False, MUTED, True), al=ALIGN_R)
    r += 2
    section(ws, f"A{r}", "فئات نسبة الحضور", "Attendance categories", 4); r += 1
    header_row(ws, r, 1, ["الفئة", "الحد الأدنى", "اسم الفئة (تلقائي)", "الوصف"], height=24); r += 1
    c0 = r
    put(ws, f"A{r}", "حضور كامل", f=font(10), al=ALIGN_R, b=box()); put(ws, f"B{r}", "—", f=font(10, False, MUTED), al=ALIGN_C, b=box())
    put(ws, f"C{r}", "حضور كامل", f=font(10, True, GREEN_T), al=ALIGN_R, b=box()); put(ws, f"D{r}", "بدون أي غياب (بعذر أو بدون عذر)", f=font(9, False, MUTED), al=ALIGN_R, b=box()); r += 1
    put(ws, f"A{r}", "ممتاز", f=font(10), al=ALIGN_R, b=box()); input_cell(ws, f"B{r}", 0.95, nf="0%")
    put(ws, f"C{r}", f'=TEXT(B{r},"0%")&" فأكثر"', f=font(10, True, GREEN_T), al=ALIGN_R, b=box()); put(ws, f"D{r}", "نسبة حضور عند الحد أو أعلى", f=font(9, False, MUTED), al=ALIGN_R, b=box()); r += 1
    put(ws, f"A{r}", "جيد", f=font(10), al=ALIGN_R, b=box()); input_cell(ws, f"B{r}", 0.90, nf="0%")
    put(ws, f"C{r}", f'=TEXT(B{r},"0%")&" – "&TEXT(B{r-1}-0.001,"0.0%")', f=font(10, True, YEL_T), al=ALIGN_R, b=box()); put(ws, f"D{r}", "بين الحدّين", f=font(9, False, MUTED), al=ALIGN_R, b=box()); r += 1
    put(ws, f"A{r}", "يحتاج دعمًا", f=font(10), al=ALIGN_R, b=box()); put(ws, f"B{r}", "—", f=font(10, False, MUTED), al=ALIGN_C, b=box())
    put(ws, f"C{r}", f'="أقل من "&TEXT(B{r-1},"0%")', f=font(10, True, RED_T), al=ALIGN_R, b=box()); put(ws, f"D{r}", "أقل من حد «جيد»", f=font(9, False, MUTED), al=ALIGN_R, b=box()); r += 1
    names["Cat_95"] = f"{cq(ws.title)}$B${c0+1}"; names["Cat_90"] = f"{cq(ws.title)}$B${c0+2}"
    names["Cat_L1"] = f"{cq(ws.title)}$C${c0}"; names["Cat_L2"] = f"{cq(ws.title)}$C${c0+1}"
    names["Cat_L3"] = f"{cq(ws.title)}$C${c0+2}"; names["Cat_L4"] = f"{cq(ws.title)}$C${c0+3}"
    names["CatLabels"] = f"{cq(ws.title)}$C${c0}:$C${c0+3}"
    r += 1
    section(ws, f"A{r}", "حالات الحضور (كما تُكتب في سجل الصف)", "Attendance statuses", 4); r += 1
    header_row(ws, r, 1, ["الحالة", "الرمز الداخلي", "يُحتسب حضورًا", "الرمز المختصر"], height=24); r += 1
    s0 = r
    for lab, code, pres, sym in [("غياب بعذر", "E", "لا", "○"), ("غياب بدون عذر", "A", "لا", "✕"), ("حاضر", "P", "نعم", "✓"), ("متأخر", "L", "نعم", "م")]:
        put(ws, f"A{r}", lab, f=font(10, True), al=ALIGN_C, b=box()); put(ws, f"B{r}", code, f=font(10), al=ALIGN_C, b=box())
        put(ws, f"C{r}", pres, f=font(10), al=ALIGN_C, b=box()); put(ws, f"D{r}", sym, f=font(10), al=ALIGN_C, b=box()); r += 1
    names["StatusLabels"] = f"{cq(ws.title)}$A${s0}:$A${s0+3}"; names["StatusCodes"] = f"{cq(ws.title)}$B${s0}:$B${s0+3}"
    names["StatusSymbols"] = f"{cq(ws.title)}$D${s0}:$D${s0+3}"
    put(ws, f"A{r}", "الخلية الفارغة في أي يوم ضمن «الحصر مكتمل حتى» تُعتبر حضورًا؛ تُختار الحالات الأخرى من القائمة.", f=font(9, False, MUTED, True), al=ALIGN_R)
    r += 2
    section(ws, f"A{r}", "الصفوف", "Classes", 4); r += 1
    header_row(ws, r, 1, ["الصف", "مفعّل؟", "اسم الورقة", "ملاحظة"], height=24); r += 1
    cl0 = r
    for i, cname in enumerate(CLASSES):
        put(ws, f"A{r}", cname, f=font(10, True), al=ALIGN_C, b=box())
        input_cell(ws, f"B{r}", "نعم" if i < ACTIVE_CLASSES else "لا")
        put(ws, f"C{r}", cname, f=font(10, False, MUTED), al=ALIGN_C, b=box())
        put(ws, f"D{r}", "" if i < ACTIVE_CLASSES else "غير مفعّل – لا يظهر في اللوحة والتقارير", f=font(9, False, MUTED), al=ALIGN_RW, b=box())
        # filter list in column G
        put(ws, f"G{cl0 + 1 + i}", f'=IF(B{r}="نعم",A{r},"")', f=font(10), al=ALIGN_C)
        r += 1
    put(ws, f"G{cl0 - 1}", "قائمة فلتر الصفوف", f=font(9, True, MUTED), al=ALIGN_C)
    put(ws, f"G{cl0}", "الكل", f=font(10), al=ALIGN_C)
    names["ClassList"] = f"{cq(ws.title)}$A${cl0}:$A${cl0+NCLS-1}"
    names["ClassActive"] = f"{cq(ws.title)}$B${cl0}:$B${cl0+NCLS-1}"
    names["FilterClassList"] = f"{cq(ws.title)}$G${cl0}:$G${cl0+NCLS}"
    names["ActiveClassList"] = f"{cq(ws.title)}$G${cl0+1}:$G${cl0+NCLS}"
    put(ws, f"A{r}", "اسم الصف يُكتب أيضًا في الخلية الصفراء أعلى ورقة الصف ويجب أن يطابق هذه القائمة.", f=font(9, False, MUTED, True), al=ALIGN_R)
    r += 2
    section(ws, f"A{r}", "طرق التواصل مع ولي الأمر", "Contact methods", 4); r += 1
    m0 = r
    for m in ["اتصال هاتفي", "رسالة نصية", "WhatsApp", "اجتماع ولي أمر", "أخرى"]:
        input_cell(ws, f"A{r}", m, al=ALIGN_R); r += 1
    names["ContactMethods"] = f"{cq(ws.title)}$A${m0}:$A${m0+4}"
    r += 1
    section(ws, f"A{r}", "الحالة الدراسية للطفل", "Student status", 4); r += 1
    st0 = r
    for a, e in [("نشط", "Active – يدخل في جميع الإحصائيات"), ("منقول", "Transferred – لا يدخل في الإحصائيات"), ("منسحب", "Withdrawn – لا يدخل في الإحصائيات")]:
        put(ws, f"A{r}", a, f=font(10, True), al=ALIGN_C, b=box()); put(ws, f"B{r}", e, f=font(9, False, MUTED), al=ALIGN_R); r += 1
    names["StudentStatusList"] = f"{cq(ws.title)}$A${st0}:$A${st0+2}"
    r += 1
    section(ws, f"A{r}", "الفصول الدراسية", "Terms", 4); r += 1
    header_row(ws, r, 1, ["الفصل", "من", "إلى"], height=24); r += 1
    t0 = r
    for name, a, b in TERMS:
        input_cell(ws, f"A{r}", name); input_cell(ws, f"B{r}", a, nf="dd/mm/yyyy"); input_cell(ws, f"C{r}", b, nf="dd/mm/yyyy"); r += 1
    names["TermNames"] = f"{cq(ws.title)}$A${t0}:$A${t0+2}"; names["TermStart"] = f"{cq(ws.title)}$B${t0}:$B${t0+2}"; names["TermEnd"] = f"{cq(ws.title)}$C${t0}:$C${t0+2}"
    r += 1
    section(ws, f"A{r}", "العطلات والإجازات الرسمية (لا تُحتسب أيامًا دراسية)", "Holidays", 4); r += 1
    header_row(ws, r, 1, ["من", "إلى", "الوصف"], height=24); r += 1
    h0 = r
    for i in range(40):
        a = HOLIDAYS[i] if i < len(HOLIDAYS) else (None, None, None)
        input_cell(ws, f"A{r}", a[0], nf="dd/mm/yyyy"); input_cell(ws, f"B{r}", a[1], nf="dd/mm/yyyy"); input_cell(ws, f"C{r}", a[2], al=ALIGN_R); r += 1
    names["HolFrom"] = f"{cq(ws.title)}$A${h0}:$A${h0+39}"; names["HolTo"] = f"{cq(ws.title)}$B${h0}:$B${h0+39}"
    # period list & yes/no & confirm list (column G, top)
    put(ws, "G4", "قوائم مساعدة", f=font(9, True, MUTED), al=ALIGN_C)
    for i, p in enumerate(["اليوم", "الأسبوع", "الشهر", "الفصل الدراسي", "العام الدراسي"]):
        put(ws, f"G{5+i}", p, f=font(10), al=ALIGN_C)
    names["PeriodList"] = f"{cq(ws.title)}$G$5:$G$9"
    put(ws, "G11", "نعم", f=font(10), al=ALIGN_C); put(ws, "G12", "لا", f=font(10), al=ALIGN_C)
    names["YesNo"] = f"{cq(ws.title)}$G$11:$G$12"
    put(ws, "G14", "✓", f=font(10), al=ALIGN_C); names["ConfirmList"] = f"{cq(ws.title)}$G$14:$G$14"
    for i, t in enumerate(["يوم", "الأسبوع", "الشهر", "الفصل الدراسي", "العام الدراسي"]):
        put(ws, f"G{16+i}", t, f=font(10), al=ALIGN_C)
    names["RptPeriodList"] = f"{cq(ws.title)}$G$16:$G$20"
    for k, v in names.items():
        if not k.startswith("_"):
            define(wb, k, v)
    # validations
    dv = DataValidation(type="list", formula1="=YesNo", allow_blank=False); ws.add_data_validation(dv)
    for c in names["_yesno_cells"]: dv.add(ws[c])
    dv2 = DataValidation(type="list", formula1="=TermNames", allow_blank=False); ws.add_data_validation(dv2); dv2.add(ws[names["_term_cell"]])
    dv3 = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=False, error="أدخل عددًا صحيحًا أكبر من صفر", errorTitle="قيمة غير صحيحة"); ws.add_data_validation(dv3)
    for i in range(1, 5): dv3.add(ws[f"B{lv0+i}"])
    dv4 = DataValidation(type="list", formula1="=YesNo", allow_blank=False); ws.add_data_validation(dv4)
    for i in range(NCLS): dv4.add(ws[f"B{cl0+i}"])
    dv5 = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)", allow_blank=True, error="أدخل تاريخًا صحيحًا", errorTitle="تاريخ غير صحيح"); ws.add_data_validation(dv5)
    dv5.add(f"A{h0}:B{h0+39}"); dv5.add(f"B{t0}:C{t0+2}")
    ws.freeze_panes = "A5"
    protect(ws)
    return names


def build_calendar(wb):
    ws = wb.create_sheet(SHEETS["cal"])
    setup_sheet(ws, tab="9CA3AF")
    title_block(ws, "التقويم الدراسي", "يُبنى تلقائيًا من الإعدادات (العطلات، الفصول) ومن سجلات الصفوف — لا يحتاج إلى إدخال", "Academic calendar (auto)")
    hdr = ["التاريخ", "اليوم", "الشهر", "مفتاح الشهر", "الأسبوع الدراسي", "الفصل", "عطلة؟", "يوم دراسي", "المسجَّلون", "حاضر", "غياب بدون عذر", "غياب بعذر", "متأخر", "نسبة الحضور", "تم التسجيل؟", "ترتيب اليوم المسجَّل"]
    header_row(ws, 4, 1, hdr, widths=[12, 10, 14, 10, 9, 12, 8, 8, 10, 8, 8, 8, 8, 10, 10, 10])
    for i, d in enumerate(DATES):
        r = CAL_FIRST + i
        put(ws, f"A{r}", d, nf="dd/mm/yyyy", al=ALIGN_C, f=font(10))
        put(ws, f"B{r}", AR_DAYS[d.weekday()], al=ALIGN_C, f=font(10))
        put(ws, f"C{r}", f"{AR_MONTHS[d.month-1]} {d.year}", al=ALIGN_C, f=font(10))
        put(ws, f"D{r}", d.year * 100 + d.month, al=ALIGN_C, f=font(10, False, MUTED))
        put(ws, f"E{r}", (d - YEAR_START).days // 7 + 1, al=ALIGN_C, f=font(10, False, MUTED))
        put(ws, f"F{r}", f'=IFERROR(IF(A{r}<=INDEX(TermEnd,MATCH(A{r},TermStart,1)),INDEX(TermNames,MATCH(A{r},TermStart,1)),"إجازة"),"")', al=ALIGN_C, f=font(10))
        put(ws, f"G{r}", f'=IF(COUNTIFS(HolFrom,"<="&A{r},HolTo,">="&A{r})>0,"عطلة","")', al=ALIGN_C, f=font(10, False, RED_T))
        put(ws, f"H{r}", f'=IF(OR(G{r}="عطلة",AND(WEEKDAY(A{r},2)=5,FridaySchool="لا")),0,1)', al=ALIGN_C, f=font(10))
        col = L(DATE_COL0 + i)
        for j, key in enumerate(["reg", "pres", "abs", "exc", "late"]):
            terms = "+".join(f"{cq(cn)}{col}{SUM_ROW[key]}" for cn in CLASSES)
            put(ws, (r, 9 + j), f"={terms}", al=ALIGN_C, f=font(10))
        put(ws, f"N{r}", f'=IF(I{r}=0,"",(J{r}+IF(LateAsPresent="نعم",M{r},0))/I{r})', al=ALIGN_C, f=font(10), nf="0.0%")
        put(ws, f"O{r}", f'=IF(I{r}>0,"نعم","")', al=ALIGN_C, f=font(10, False, GREEN_T))
        put(ws, f"P{r}", f'=IF(I{r}>0,COUNTIF($I${CAL_FIRST}:I{r},">0"),"")', al=ALIGN_C, f=font(10, False, MUTED))
    names = {"Cal_Date": "A", "Cal_Day": "B", "Cal_MonthLabel": "C", "Cal_MonthKey": "D", "Cal_Week": "E", "Cal_Term": "F", "Cal_Hol": "G", "Cal_School": "H",
             "Cal_Reg": "I", "Cal_Pres": "J", "Cal_Abs": "K", "Cal_Exc": "L", "Cal_Late": "M", "Cal_Rate": "N", "Cal_Recorded": "O", "Cal_RecIdx": "P"}
    for k, c in names.items():
        define(wb, k, f"{cq(ws.title)}${c}${CAL_FIRST}:${c}${CAL_LAST}")
    tab = Table(displayName="tblCalendar", ref=f"A4:P{CAL_LAST}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A5"
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(f"A{CAL_FIRST}:P{CAL_LAST}", FormulaRule(formula=[f"$H{CAL_FIRST}=0"], fill=fill(GRAY_F), font=Font(name=FONT, color=MUTED)))
    ws.conditional_formatting.add(f"A{CAL_FIRST}:A{CAL_LAST}", FormulaRule(formula=[f"$A{CAL_FIRST}=TODAY()"], fill=fill(BLUE_F), font=Font(name=FONT, bold=True)))
    protect(ws)
    return ws
