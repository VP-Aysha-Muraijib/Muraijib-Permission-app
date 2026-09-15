# -*- coding: utf-8 -*-
"""Independent Python model of the attendance rules; compares against LibreOffice-recalculated values."""
import sys, datetime as dt
from openpyxl import load_workbook
from common import *
import testdata

path = sys.argv[1]
wb = load_workbook(path, data_only=True)
students = testdata.make_students()
logs = testdata.make_contacts(students)
T3, T5, T10, T15 = 3, 5, 10, 15
fails = []
def check(label, got, exp):
    ok = (got == exp) if not isinstance(exp, float) else (isinstance(got, (int, float)) and abs(got - exp) < 1e-6)
    if not ok:
        fails.append(f"{label}: got {got!r} expected {exp!r}")

def to_date(v):
    if isinstance(v, dt.datetime): return v.date()
    return v

# ---- model
model = {}   # id -> dict
for ci, cn in enumerate(CLASSES[:ACTIVE_CLASSES]):
    for i, s in enumerate(students[cn]):
        active = True
        join = YEAR_START
        eff = [None] * NDAYS
        for d in range(NDAYS):
            if not active or DATES[d] < join: continue
            v = s["grid"][d]
            if v: eff[d] = v
            elif testdata.confirmed(ci, d): eff[d] = "حاضر"
        P = eff.count("حاضر"); A = eff.count("غائب"); E = eff.count("بعذر"); Lt = eff.count("متأخر")
        tot = P + A + E + Lt
        rate = (P + Lt) / tot if tot else None
        level = 15 if A >= 15 else 10 if A >= 10 else 5 if A >= 5 else 3 if A >= 3 else 0
        absdates = [DATES[d] for d in range(NDAYS) if eff[d] == "غائب"]
        reached = {n: (absdates[n - 1] if len(absdates) >= n else None) for n in (3, 5, 10, 15)}
        lastc = max([l[0] for l in logs if l[1] == s["id"]], default=None)
        contacted = None if not active else ("—" if level == 0 else ("نعم" if (lastc and lastc >= reached[level]) else "لا"))
        model[s["id"]] = dict(cls=cn, ci=ci, slot=i, active=active, P=P, A=A, E=E, L=Lt, tot=tot, rate=rate, level=level, reached=reached,
                              absdates=absdates, lastc=lastc, contacted=contacted, name=s["name"], events=sum(1 for v in eff if v in ("غائب", "بعذر")),
                              status_today=eff[testdata.TODAY_IDX], today_level=next((n for n in (15, 10, 5, 3) if reached[n] == DATES[testdata.TODAY_IDX]), 0))

# ---- registers
for cn in CLASSES[:ACTIVE_CLASSES]:
    ws = wb[cn]
    for i, s in enumerate(students[cn]):
        r = REG_FIRST + i; m = model[s["id"]]
        check(f"{cn} r{r} present", ws[f"N{r}"].value, m["P"]); check(f"{cn} r{r} absent", ws[f"D{r}"].value, m["A"])
        check(f"{cn} r{r} excused", ws[f"E{r}"].value, m["E"]); check(f"{cn} r{r} late", ws[f"F{r}"].value, m["L"])
        check(f"{cn} r{r} days", ws[f"M{r}"].value, m["tot"]); check(f"{cn} r{r} rate", ws[f"G{r}"].value, m["rate"])
        check(f"{cn} r{r} level", ws[f"T{r}"].value, m["level"])
        for col, n in zip(["V", "W", "X", "Y"], (3, 5, 10, 15)):
            check(f"{cn} r{r} reached{n}", to_date(ws[f"{col}{r}"].value), m["reached"][n])
        check(f"{cn} r{r} lastcontact", to_date(ws[f"R{r}"].value), m["lastc"])
        check(f"{cn} r{r} contacted", ws[f"Q{r}"].value, m["contacted"])
        exp_status = {0: "منتظم", 3: "⚠ يحتاج متابعة", 5: "⚠ غياب متكرر", 10: "⚠ غياب مرتفع", 15: "🔴 حالة حرجة"}[m["level"]]
        check(f"{cn} r{r} status", ws[f"H{r}"].value, exp_status)
        check(f"{cn} r{r} phone", ws[f"J{r}"].value, s["phone"])
# summary rows for today
ws = wb[CLASSES[0]]; col = L(DATE_COL0 + testdata.TODAY_IDX)
act = [model[s["id"]] for s in students[CLASSES[0]] if model[s["id"]]["active"]]
check("A today abs", ws[f"{col}{SUM_ROW['abs']}"].value, sum(1 for m in act if m["status_today"] == "غائب"))
check("A today reg", ws[f"{col}{SUM_ROW['reg']}"].value, sum(1 for m in act if m["status_today"]))
wsE = wb[CLASSES[4]]; check("E today reg (unconfirmed)", wsE[f"{col}{SUM_ROW['reg']}"].value, 0)
# ---- master
wm = wb[SHEETS["stu"]]
for ci, cn in enumerate(CLASSES[:ACTIVE_CLASSES]):
    for i, s in enumerate(students[cn]):
        r = STU_FIRST + ci * SLOTS + i; m = model[s["id"]]
        check(f"master r{r} id", wm[f"B{r}"].value, s["id"]); check(f"master r{r} class", wm[f"D{r}"].value, cn)
        if m["active"]:
            check(f"master r{r} abs", wm[f"L{r}"].value, m["A"]); check(f"master r{r} rate", wm[f"O{r}"].value, m["rate"])
            check(f"master r{r} today", wm[f"AG{r}"].value, {"حاضر": "P", "غائب": "A", "بعذر": "E", "متأخر": "L", None: None}[m["status_today"]])
            check(f"master r{r} lvltoday", wm[f"AJ{r}"].value, m["today_level"])
# ---- dashboard
wd = wb[SHEETS["dash"]]; wc = wb[SHEETS["calc"]]
actives = [m for m in model.values() if m["active"]]
check("dash date", to_date(wc["B3"].value), DATES[testdata.TODAY_IDX])
check("kpi total", wd["A8"].value, len(actives))
check("kpi present today", wd["C8"].value, sum(1 for m in actives if m["status_today"] in ("حاضر", "متأخر")))
check("kpi absent today", wd["E8"].value, sum(1 for m in actives if m["status_today"] in ("غائب", "بعذر")))
reg_t = sum(1 for m in actives if m["status_today"]); att_t = sum(1 for m in actives if m["status_today"] in ("حاضر", "متأخر"))
check("kpi rate today", wd["G8"].value, att_t / reg_t)
for col, n in zip("IKMO", (3, 5, 10, 15)):
    check(f"kpi >= {n}", wd[f"{col}8"].value, sum(1 for m in actives if m["A"] >= n))
# alerts today
new_today = sorted([m for m in actives if m["today_level"] > 0], key=lambda m: (m["ci"], m["slot"]))
got = [wd[f"A{18+k}"].value for k in range(10) if wd[f"A{18+k}"].value]
check("alerts today names", got, [m["name"] for m in new_today])
check("alerts levels", [wd[f"C{18+k}"].value for k in range(len(new_today))], [f"بلغ {m['today_level']} غيابات" for m in new_today])
# overall month rate (November: recorded days idx 45..48)
nov = [d for d in range(NDAYS) if DATES[d].month == 11 and DATES[d].year == 2026]
def eff_status(m, d):
    s = next(x for x in students[m["cls"]] if x["id"] == [k for k, v in model.items() if v is m][0])
    join = YEAR_START
    if DATES[d] < join: return None
    v = s["grid"][d]
    return v if v else ("حاضر" if testdata.confirmed(m["ci"], d) else None)
att = tot = 0
for m in actives:
    for d in nov:
        v = eff_status(m, d)
        if v: tot += 1; att += v in ("حاضر", "متأخر")
check("overall month rate", wd["I13"].value, att / tot)
# categories
cats = {"حضور كامل": 0, "95% فأكثر": 0, "90% – 94.9%": 0, "أقل من 90%": 0}
for m in actives:
    if m["tot"] == 0: continue
    if m["A"] == 0 and m["E"] == 0: cats["حضور كامل"] += 1
    elif m["rate"] >= 0.95: cats["95% فأكثر"] += 1
    elif m["rate"] >= 0.90: cats["90% – 94.9%"] += 1
    else: cats["أقل من 90%"] += 1
for k, (lab, n) in enumerate(cats.items()):
    check(f"category {lab}", (wd[f"I{19+k}"].value, wd[f"J{19+k}"].value), (lab, n))
# class comparison today
for k, cn in enumerate(CLASSES[:ACTIVE_CLASSES]):
    ms = [m for m in actives if m["cls"] == cn]
    r = 31 + k
    check(f"cls {cn} kids", wd[f"G{r}"].value, len(ms))
    check(f"cls {cn} present", wd[f"H{r}"].value, sum(1 for m in ms if m["status_today"] in ("حاضر", "متأخر")))
    regc = sum(1 for m in ms if m["status_today"]); attc = sum(1 for m in ms if m["status_today"] in ("حاضر", "متأخر"))
    check(f"cls {cn} rate", wd[f"J{r}"].value, (attc / regc) if regc else None)
    check(f"cls {cn} 3+", wd[f"L{r}"].value, sum(1 for m in ms if m["A"] >= 3))
check("class rows 6-7", (wd["F36"].value, wd["F37"].value), (CLASSES[5], CLASSES[6]))
# top attendance ordering & support list
top = sorted([m for m in actives if m["tot"] > 0], key=lambda m: (-round(m["rate"] * 10000), -m["P"], m["ci"], m["slot"]))[:10]
check("top10 names", [wd[f"A{62+k}"].value for k in range(10)], [("⭐ " if k < 3 else "") + m["name"] for k, m in enumerate(top)])
sup = sorted([m for m in actives if m["tot"] > 0], key=lambda m: (-(m["A"] + m["E"]), -round((1 - m["rate"]) * 10000), m["ci"], m["slot"]))[:10]
check("support10 names", [wd[f"F{62+k}"].value for k in range(10)], [m["name"] for m in sup])
crit = sorted([m for m in actives if m["A"] >= 15], key=lambda m: (-m["A"], m["ci"], m["slot"]))
check("critical names", [wd[f"L{62+k}"].value for k in range(len(crit))], [m["name"] for m in crit])
check("critical blank after", wd[f"L{62+len(crit)}"].value, None)
# follow-up sheet
wf = wb[SHEETS["fu"]]
fu = sorted([m for m in actives if m["level"] > 0], key=lambda m: (-m["A"], m["ci"], m["slot"]))
check("followup count", wf["B3"].value, len(fu))
check("followup names", [wf[f"B{6+k}"].value for k in range(len(fu))], [m["name"] for m in fu])
check("followup contacted", [wf[f"I{6+k}"].value for k in range(len(fu))], [m["contacted"] for m in fu])
# absence log
wa = wb[SHEETS["alog"]]
check("absence log total", wa["B3"].value, sum(m["events"] for m in actives))
check("absence log first row filled", wa["A6"].value is not None, True)
check("absence log row after last blank", wa[f"A{6+sum(m['events'] for m in actives)}"].value, None)
# absence log cum for a known critical student (A slot 6): find rows with that name
mA6 = model[students[CLASSES[0]][5]["id"]]
rows = [r for r in range(6, 400) if wa[f"B{r}"].value == mA6["name"] and wa[f"E{r}"].value == "غائب"]
check("alog A6 cum sequence", [wa[f"F{r}"].value for r in rows], list(range(1, mA6["A"] + 1)))
check("alog A6 dates", [to_date(wa[f"A{r}"].value) for r in rows], mA6["absdates"])
# calendar
wcal = wb[SHEETS["cal"]]
check("cal recorded days", sum(1 for r in range(CAL_FIRST, CAL_LAST + 1) if wcal[f"O{r}"].value == "نعم"), testdata.REC_DAYS)
check("cal holiday Dec 2", wcal[f"H{CAL_FIRST + DATES.index(dt.date(2026,12,2))}"].value, 0)
check("cal today registered", wcal[f"I{CAL_FIRST + testdata.TODAY_IDX}"].value, reg_t)
# profile (A slot 6 selected)
wp = wb[SHEETS["prof"]]
check("profile abs", wp["N9"].value, mA6["A"])
check("profile absdates", [to_date(wp[f"B{19+k}"].value) for k in range(min(20, mA6["A"]))], mA6["absdates"][:20])
check("profile contacts", [wp[f"N{19+k}"].value for k in range(2)], [l[2] for l in logs if l[1] == students[CLASSES[0]][5]["id"]][:2])
# daily sheet (KG1-A, last recorded day)
wdy = wb[SHEETS["daily"]]
check("daily class reg", wdy["E8"].value, sum(1 for m in actives if m["cls"] == CLASSES[0] and m["status_today"]))
check("daily class5 not counted", wdy[f"C{11+SLOTS+3+2+4}"].value, "غير محصور")
# report (month, all)
wr = wb[SHEETS["rpt"]]
check("report kids", wr["A8"].value, len(actives)); check("report rate (daily)", wr["G8"].value, att_t / reg_t)
absent_today = sorted([m for m in actives if m["status_today"] in ("غائب", "بعذر")], key=lambda m: (m["ci"], m["slot"]))
check("report absentees", [wr[f"A{74+k}"].value for k in range(len(absent_today) + 1)], [m["name"] for m in absent_today] + [None])
check("report day row", (to_date(wr["A46"].value), wr["C46"].value, wr["A47"].value), (DATES[testdata.TODAY_IDX], reg_t, None))
# trends monthly Nov row (Aug=row r1+2 ... ) : find by label
wt = wb[SHEETS["trend"]]
for r in range(1, 120):
    if wt[f"A{r}"].value == "نوفمبر 2026" and isinstance(wt[f"B{r}"].value, (int, float)):
        check("trend nov days", wt[f"B{r}"].value, 4); check("trend nov rate", wt[f"H{r}"].value, att / tot); break
else:
    fails.append("trend nov row not found")
# no error values anywhere
errs = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#") and c.value[1:4].isupper(): errs += 1
check("error cells", errs, 0)
print("checks failed:", len(fails))
for f in fails[:60]: print(" -", f)
print("model summary: active", len(actives), "| today reg/att", reg_t, att_t, "| >=3:", sum(1 for m in actives if m['A']>=3), ">=15:", sum(1 for m in actives if m['A']>=15), "| new today:", [(m['name'], m['cls'], m['today_level']) for m in new_today])
